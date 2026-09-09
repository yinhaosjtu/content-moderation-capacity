
import gc
import json
import math
import os
import statistics
import time

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from Solvers import solver_benders as sb
from Solvers import solver_heuristic as sh


_CACHE_INTEGER_STATUSES = {
    2: "ok",
    3: "infeasible",
    4: "inf_or_unbd",
    5: "unbounded",
    6: "cutoff",
    7: "iteration_limit",
    8: "node_limit",
    9: "time_limit",
    10: "solution_limit",
    11: "interrupted",
    12: "numeric",
    13: "suboptimal",
    15: "user_obj_limit",
    16: "work_limit",
    17: "oom",
}


def cache_identity(unit, instance):
    return str(unit), str(instance)


def load_cache(path):
    cache = {}
    if not os.path.exists(path):
        return cache
    with open(path, "r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                entry = json.loads(line)
            except (TypeError, ValueError):
                continue
            unit = entry.get("unit")
            instance = entry.get("instance")
            if isinstance(unit, str) and isinstance(instance, str) and "result" in entry:
                cache[cache_identity(unit, instance)] = entry["result"]
    return cache


def _cache_json_default(value):
    item = getattr(value, "item", None)
    if callable(item):
        return item()
    raise TypeError(repr(value))


def store_cache(path, unit, instance, result):
    cache_dir = os.path.dirname(path)
    if cache_dir and not os.path.exists(cache_dir):
        os.makedirs(cache_dir, exist_ok=True)

    payload = {"unit": str(unit), "instance": str(instance), "result": result}
    if os.path.exists(path) and os.path.getsize(path) > 0:
        with open(path, "rb+") as boundary:
            boundary.seek(-1, os.SEEK_END)
            if boundary.read(1) != b"\n":
                boundary.seek(0, os.SEEK_END)
                boundary.write(b"\n")
                boundary.flush()
                os.fsync(boundary.fileno())
    with open(path, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(
            payload,
            ensure_ascii=False,
            allow_nan=True,
            default=_cache_json_default,
        ) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def exception_status(error):
    text = "{0}: {1}".format(type(error).__name__, error).lower()
    oom_markers = (
        "out of memory", "out-of-memory", "memory limit", "mem_limit",
        "cannot allocate memory", "failed to allocate", "allocation failed",
        "gurobi error 10001",
    )
    if (isinstance(error, MemoryError)
            or getattr(error, "errno", None) == 10001
            or any(marker in text for marker in oom_markers)):
        return "oom"
    timeout_markers = ("time limit", "time-limit", "timed out", "timeout")
    if isinstance(error, TimeoutError) or any(
            marker in text for marker in timeout_markers):
        return "time_limit"
    return "error"


def error_text(error):
    return "{0}: {1}".format(type(error).__name__, error)


def normalized_solver_status(result, time_limit=None, time_field="time"):
    if not isinstance(result, dict):
        return "error"
    raw = result.get("status")
    status = None
    if isinstance(raw, bool):
        status = None
    elif isinstance(raw, (int, float)) and math.isfinite(float(raw)):
        status = _CACHE_INTEGER_STATUSES.get(
            int(raw), "status_{0}".format(int(raw)))
    elif raw is not None:
        status = str(raw).strip().lower().replace("-", "_").replace(" ", "_")
        aliases = {
            "optimal": "ok",
            "loaded": "ok",
            "timelimit": "time_limit",
            "memory_limit": "oom",
            "out_of_memory": "oom",
            "mem_limit": "oom",
        }
        status = aliases.get(status, status)
    elapsed = result.get(time_field)
    try:
        elapsed = float(elapsed)
    except (TypeError, ValueError):
        elapsed = None
    if status is None and time_limit is not None and elapsed is not None:
        threshold = max(
            float(time_limit) - max(1.0, 0.001 * float(time_limit)), 0.0)
        if elapsed >= threshold:
            status = "time_limit"
    if status is None:
        status = "ok" if result.get("objective") is not None else "no_solution"
    return status


def is_cacheable_result(result):
    return isinstance(result, dict) and result.get("_status", "ok") != "error"

INSTANCE_DIR = "instances"
RESULT_FILE = "results_ablation.xlsx"
CACHE_FILE = "cache/cache_ablation_experiment.jsonl"
HEURISTIC_CACHE_FILE = "cache/cache_heuristic_ablation.jsonl"
COMPARATIVE_CACHE_FILE = "cache/cache_comparative_experiment.jsonl"


def discover_ordinals(directory, scale):
    ordinals = []
    prefix = scale + "_"
    for name in os.listdir(directory):
        if not (name.startswith(prefix) and name.endswith(".json")):
            continue
        stem = os.path.splitext(name)[0]
        parts = stem.rsplit("_", 1)
        if len(parts) == 2 and parts[1].isdigit():
            ordinals.append(int(parts[1]))
    return sorted(set(ordinals))


TIME_LIMIT = 3600.0
MIP_GAP = 0.0
THREADS = 0
HEURISTIC_TIME_LIMIT = 300.0
REFERENCE_REL_TOL = 1e-5
REFERENCE_GAP_TOL_PERCENT = 1e-3
BENDERS_ABLATION_INSTANCE_LIMITS = {"Medium": 10, "Small": 10}

N_DAYS_MEDIUM = 10

MEDIUM_CONFIGS = [
    ("full_BBC", {}),
    ("block_Bhalf", {"recourse_parts": N_DAYS_MEDIUM // 2,
                     "adaptive_block_rounds": 0}),
    ("warmstart_off", {"mip_start": False}),
    ("incumbent_cuts_off", {"incumbent_cuts": False}),
    ("root_node_cuts_off", {"node_cut_rounds": 0}),
    ("shallow_tree_cuts_off", {"tree_cut_rounds": 0}),
    ("all_mipnode_cuts_off", {"node_cut_rounds": 0,
                              "tree_cut_rounds": 0}),
    ("final_purge_off", {"final_cut_keep": -1}),
]

PANELS = [
    ("medium_ablation", "Medium", "M", MEDIUM_CONFIGS),
]

HEURISTIC_BASE_PARAMS = {
    "time_limit": HEURISTIC_TIME_LIMIT,
    "threads": THREADS,
    "output": False,
    "root_rounds": 10,
    "integer_rounds": 6,
    "local_rounds": 2,
    "max_neighbors": 12,
    "search_scenarios": 0,
    "master_gap": 0.01,
    "quality_gap": 0.005,
    "adaptive": True,
}

HEURISTIC_CONFIGS = [
    (
        "H0",
        "H0_base_construction_v1",
        {
            "enable_multistart": False,
            "enable_integer_stage": False,
            "enable_local_search": False,
        },
    ),
    (
        "H1",
        "H1_diversified_multistart_v1",
        {
            "enable_multistart": True,
            "enable_integer_stage": False,
            "enable_local_search": False,
        },
    ),
    (
        "H2",
        "H2_integer_refinement_v1",
        {
            "enable_multistart": True,
            "enable_integer_stage": True,
            "enable_local_search": False,
        },
    ),
    (
        "H3",
        "H3_full_heuristic_v1",
        {
            "enable_multistart": True,
            "enable_integer_stage": True,
            "enable_local_search": True,
        },
    ),
]

YELLOW = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00",
                     fill_type="solid")


def instance_paths(scale, label_prefix):
    out = []
    ordinals = discover_ordinals(INSTANCE_DIR, scale)
    for i, ordinal in enumerate(ordinals, start=1):
        path = os.path.join(INSTANCE_DIR, "%s_%02d.json" % (scale, ordinal))
        out.append(("%s%d" % (label_prefix, i), path))
    return out


def build_columns(configs):
    cols = ["Statistic"]
    widths = [12]
    for cfg_name, _ in configs:
        cols.append("%s_Gap(%%)" % cfg_name)
        cols.append("%s_Time(s)" % cfg_name)
        widths.append(16)
        widths.append(16)
    return cols, widths


def render(value, sci=False):
    if value is None:
        return "-"
    if sci:
        return "%.3e" % value
    return "%.2f" % value


def cap_report_time(value, time_limit=TIME_LIMIT):
    if value is None:
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(value):
        return None
    return min(max(value, 0.0), float(time_limit))


def json_scalar(value):
    if value is None:
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) else None


def complete_mean(values, expected_count):
    if expected_count <= 0 or len(values) != expected_count:
        return None
    normalized = [json_scalar(value) for value in values]
    if any(value is None for value in normalized):
        return None
    return statistics.fmean(normalized)


def slim_result(result, status="ok", error=None):
    result = result or {}
    return {
        "objective": json_scalar(result.get("objective")),
        "bound": json_scalar(result.get("bound")),
        "gap": json_scalar(result.get("gap")),
        "time": json_scalar(result.get("time")),
        "_status": status,
        "_error": error,
    }


def solve_config(instance, params):
    started = time.perf_counter()
    try:
        raw = sb.solve_instance(instance, params)
    except Exception as exc:
        return slim_result(
            {"time": time.perf_counter() - started},
            status=exception_status(exc),
            error=error_text(exc),
        )
    status = normalized_solver_status(
        raw, time_limit=params.get("time_limit"), time_field="time")
    return slim_result(raw, status=status)


def slim_heuristic_result(result, status="ok", error=None):
    compact = slim_result(result, status=status, error=error)
    result = result or {}
    raw_stats = result.get("stats")
    if not isinstance(raw_stats, dict):
        raw_stats = {}
    fields = (
        "enable_multistart",
        "enable_integer_stage",
        "enable_local_search",
        "root_time",
        "integer_time",
        "local_time",
        "root_rounds",
        "integer_rounds",
        "local_iterations",
        "evaluations",
        "lp_solves",
        "candidate_count",
        "pricing_time",
        "search_scenarios",
        "total_scenarios",
        "cut_batch_size",
        "workers",
        "lp_backend",
    )
    diagnostics = {}
    for field in fields:
        value = raw_stats.get(field)
        if isinstance(value, bool):
            diagnostics[field] = value
        elif isinstance(value, str):
            diagnostics[field] = value
        else:
            scalar = json_scalar(value)
            if scalar is not None:
                diagnostics[field] = scalar
    compact["iterations"] = json_scalar(result.get("iterations"))
    compact["stats"] = diagnostics
    return compact


def solve_heuristic_config(instance, params):
    started = time.perf_counter()
    try:
        raw = sh.solve_instance(instance, params)
    except Exception as exc:
        return slim_heuristic_result(
            {"time": time.perf_counter() - started},
            status=exception_status(exc),
            error=error_text(exc),
        )
    status = normalized_solver_status(
        raw, time_limit=params.get("time_limit"), time_field="time"
    )
    return slim_heuristic_result(raw, status=status)


def validated_reference_objective(cache, instance_name):
    objectives = []
    for method in ("benders", "lshaped"):
        result = cache.get(cache_identity(method, instance_name))
        if not isinstance(result, dict):
            raise RuntimeError(
                "Missing {0} reference for {1}".format(
                    method, instance_name
                )
            )
        status = str(result.get("_status", "ok")).strip().lower()
        if status != "ok":
            raise RuntimeError(
                "Invalid {0} reference status for {1}: {2}".format(
                    method, instance_name, status
                )
            )
        objective = json_scalar(result.get("objective"))
        gap = json_scalar(result.get("gap"))
        if objective is None or gap is None:
            raise RuntimeError(
                "Incomplete {0} reference for {1}".format(
                    method, instance_name
                )
            )
        if abs(gap) > REFERENCE_GAP_TOL_PERCENT:
            raise RuntimeError(
                "Inexact {0} reference for {1}: gap={2}".format(
                    method, instance_name, gap
                )
            )
        objectives.append((method, objective))
    scale = max(
        1.0, abs(objectives[0][1]), abs(objectives[1][1])
    )
    difference = abs(objectives[0][1] - objectives[1][1])
    if difference > REFERENCE_REL_TOL * scale:
        raise RuntimeError(
            "Reference objectives disagree for {0}: "
            "benders={1}, lshaped={2}".format(
                instance_name, objectives[0][1], objectives[1][1]
            )
        )
    return min(objective for _, objective in objectives)


def relative_deviation(objective, reference, instance_name, config_name):
    objective = json_scalar(objective)
    reference = json_scalar(reference)
    if objective is None:
        return None
    if reference is None or abs(reference) <= 1e-12:
        raise RuntimeError(
            "Invalid reference objective for {0}".format(instance_name)
        )
    value = 100.0 * (objective - reference) / abs(reference)
    if value < -100.0 * REFERENCE_REL_TOL:
        raise RuntimeError(
            "Negative RPD for {0}/{1}: {2}".format(
                instance_name, config_name, value
            )
        )
    return value


def prepare_sheet(book, sheet_name, configs, first):
    cols, widths = build_columns(configs)
    sheet = book.active if first else book.create_sheet()
    sheet.title = sheet_name
    sheet.append(cols)
    for idx, w in enumerate(widths):
        sheet.column_dimensions[get_column_letter(idx + 1)].width = w + 2
    time_cols = [i + 1 for i, c in enumerate(cols) if c.endswith("_Time(s)")]
    for c in time_cols:
        sheet.cell(row=1, column=c).fill = YELLOW
    return sheet, cols, widths, time_cols


def append_capped_row(
    sheet, record, time_cols, time_limit=TIME_LIMIT
):
    record = list(record)
    for column in time_cols:
        record[column - 1] = cap_report_time(
            record[column - 1], time_limit
        )
    sheet.append(record)
    row = sheet.max_row
    for column in time_cols:
        sheet.cell(row=row, column=column).fill = YELLOW


def run_panel(book, sheet_name, scale, label_prefix, configs, first, cache):
    sheet, cols, widths, time_cols = prepare_sheet(
        book, sheet_name, configs, first)
    book.save(RESULT_FILE)

    print("[%s]" % sheet_name, flush=True)
    paths = instance_paths(scale, label_prefix)[
        :BENDERS_ABLATION_INSTANCE_LIMITS[scale]
    ]
    aggregates = {
        cfg_name: {"gap": [], "time": []} for cfg_name, _ in configs
    }
    for label, path in paths:
        with open(path, "r", encoding="utf-8") as h:
            inst = json.load(h)
        instance_name = inst.get("meta", {}).get(
            "name", os.path.splitext(os.path.basename(path))[0])
        for cfg_name, delta in configs:
            params = {"time_limit": TIME_LIMIT, "mip_gap": MIP_GAP,
                      "threads": THREADS}
            params.update(delta)
            unit = "{0}/{1}".format(sheet_name, cfg_name)
            key = cache_identity(unit, instance_name)
            if key in cache:
                r = cache[key]
                tag = "cache"
            else:
                r = solve_config(inst, params)
                if not is_cacheable_result(r):
                    tag = "error-not-cached"
                else:
                    store_cache(CACHE_FILE, unit, instance_name, r)
                    cache[key] = r
                    tag = "terminal-cached" if r.get(
                        "_status") in ("time_limit", "oom") else "solved"
            g = r.get("gap")
            t = cap_report_time(r.get("time"))
            aggregates[cfg_name]["gap"].append(json_scalar(g))
            aggregates[cfg_name]["time"].append(t)
            print("[%s] %s/%s: status=%s obj=%s (%s)"
                  % (sheet_name, label, cfg_name,
                     r.get("_status", "unknown"), r.get("objective"), tag),
                  flush=True)
            gc.collect()
        del inst
        gc.collect()
    record = ["Mean"]
    for cfg_name, _ in configs:
        record.extend(
            [
                complete_mean(aggregates[cfg_name]["gap"], len(paths)),
                complete_mean(aggregates[cfg_name]["time"], len(paths)),
            ]
        )
    append_capped_row(sheet, record, time_cols)
    book.save(RESULT_FILE)


def run_heuristic_stages(
    book, heuristic_cache, reference_cache
):
    summary = {}
    for scale, label_prefix in (("Small", "S"), ("Medium", "M")):
        for display_name, _, _ in HEURISTIC_CONFIGS:
            summary[(scale, display_name)] = []
        for ordinal, (_, path) in enumerate(
            instance_paths(scale, label_prefix), start=1
        ):
            with open(path, "r", encoding="utf-8") as handle:
                instance = json.load(handle)
            instance_name = instance.get("meta", {}).get(
                "name", os.path.splitext(os.path.basename(path))[0]
            )
            reference = validated_reference_objective(
                reference_cache, instance_name
            )
            for display_name, unit_name, delta in HEURISTIC_CONFIGS:
                params = dict(HEURISTIC_BASE_PARAMS)
                params.update(delta)
                unit = "heuristic_stages/{0}".format(unit_name)
                key = cache_identity(unit, instance_name)
                if key in heuristic_cache:
                    result = heuristic_cache[key]
                    tag = "cache"
                else:
                    result = solve_heuristic_config(instance, params)
                    if is_cacheable_result(result):
                        store_cache(
                            HEURISTIC_CACHE_FILE,
                            unit,
                            instance_name,
                            result,
                        )
                        heuristic_cache[key] = result
                        tag = (
                            "terminal-cached"
                            if result.get("_status")
                            in ("time_limit", "oom")
                            else "solved"
                        )
                    else:
                        tag = "error-not-cached"
                objective = json_scalar(result.get("objective"))
                rpd = relative_deviation(
                    objective,
                    reference,
                    instance_name,
                    display_name,
                )
                elapsed = cap_report_time(
                    result.get("time"), HEURISTIC_TIME_LIMIT
                )
                summary[(scale, display_name)].append(
                    (rpd, elapsed)
                    if rpd is not None and elapsed is not None
                    else None
                )
                print(
                    "[heuristic_stages] {0}{1}/{2}: "
                    "status={3} obj={4} rpd={5} ({6})".format(
                        label_prefix,
                        ordinal,
                        display_name,
                        result.get("_status", "unknown"),
                        objective,
                        rpd,
                        tag,
                    ),
                    flush=True,
                )
                gc.collect()
            del instance
            gc.collect()
    return summary


def write_heuristic_summary(book, summary):
    sheet = book.create_sheet("heuristic_ablation")
    columns = [
        "Scale",
        "Configuration",
        "Mean_RPD(%)",
        "Mean_Time(s)",
    ]
    sheet.append(columns)
    widths = [10, 20, 16, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.cell(row=1, column=4).fill = YELLOW
    for scale in ("Small", "Medium"):
        for display_name, _, _ in HEURISTIC_CONFIGS:
            values = summary[(scale, display_name)]
            if values and all(value is not None for value in values):
                rpds = [value[0] for value in values]
                times = [value[1] for value in values]
                row = [
                    scale,
                    display_name,
                    statistics.fmean(rpds),
                    cap_report_time(
                        statistics.fmean(times), HEURISTIC_TIME_LIMIT
                    ),
                ]
            else:
                row = [scale, display_name, None, None]
            sheet.append(row)
            sheet.cell(row=sheet.max_row, column=4).fill = YELLOW


def main():
    cache = load_cache(CACHE_FILE)
    heuristic_cache = load_cache(HEURISTIC_CACHE_FILE)
    reference_cache = load_cache(COMPARATIVE_CACHE_FILE)
    book = Workbook()
    print("[cache] loaded {0} ablation units from {1}".format(
        len(cache), CACHE_FILE), flush=True)
    for index, (sheet_name, scale, label_prefix, configs) in enumerate(PANELS):
        run_panel(book, sheet_name, scale, label_prefix, configs,
                  first=(index == 0), cache=cache)
    print("[cache] loaded {0} heuristic units from {1}".format(
        len(heuristic_cache), HEURISTIC_CACHE_FILE), flush=True)
    summary = run_heuristic_stages(
        book, heuristic_cache, reference_cache
    )
    write_heuristic_summary(book, summary)
    book.save(RESULT_FILE)


if __name__ == "__main__":
    main()
