
import gc
import importlib
import json
import math
import os
import re
import time

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

_CACHE_INTEGER_STATUSES = {
    2: "ok",
    3: "infeasible",
    4: "inf_or_unbd",
    5: "unbounded",
    6: "cutoff",
    7: "iteration_limit",
    8: "node_limit",
    9: "time_limit",
    10: "solution_limit",
    11: "interrupted",
    12: "numeric",
    13: "suboptimal",
    15: "user_obj_limit",
    16: "work_limit",
    17: "oom",
}


def cache_identity(unit, instance):
    return str(unit), str(instance)


def load_cache(path):
    cache = {}
    if not os.path.exists(path):
        return cache
    with open(path, "r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                entry = json.loads(line)
            except (TypeError, ValueError):
                continue
            unit = entry.get("unit")
            instance = entry.get("instance")
            if isinstance(unit, str) and isinstance(instance, str) and "result" in entry:
                cache[cache_identity(unit, instance)] = entry["result"]
    return cache


def _cache_json_default(value):
    item = getattr(value, "item", None)
    if callable(item):
        return item()
    raise TypeError(repr(value))


def store_cache(path, unit, instance, result):
    cache_dir = os.path.dirname(path)
    if cache_dir and not os.path.exists(cache_dir):
        os.makedirs(cache_dir, exist_ok=True)

    payload = {"unit": str(unit), "instance": str(instance), "result": result}
    if os.path.exists(path) and os.path.getsize(path) > 0:
        with open(path, "rb+") as boundary:
            boundary.seek(-1, os.SEEK_END)
            if boundary.read(1) != b"\n":
                boundary.seek(0, os.SEEK_END)
                boundary.write(b"\n")
                boundary.flush()
                os.fsync(boundary.fileno())
    with open(path, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(
            payload,
            ensure_ascii=False,
            allow_nan=True,
            default=_cache_json_default,
        ) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def exception_status(error):
    text = "{0}: {1}".format(type(error).__name__, error).lower()
    oom_markers = (
        "out of memory", "out-of-memory", "memory limit", "mem_limit",
        "cannot allocate memory", "failed to allocate", "allocation failed",
        "gurobi error 10001",
    )
    if (isinstance(error, MemoryError)
            or getattr(error, "errno", None) == 10001
            or any(marker in text for marker in oom_markers)):
        return "oom"
    timeout_markers = ("time limit", "time-limit", "timed out", "timeout")
    if isinstance(error, TimeoutError) or any(
            marker in text for marker in timeout_markers):
        return "time_limit"
    return "error"


def error_text(error):
    return "{0}: {1}".format(type(error).__name__, error)


def normalized_solver_status(result, time_limit=None, time_field="time"):
    if not isinstance(result, dict):
        return "error"
    raw = result.get("status")
    status = None
    if isinstance(raw, bool):
        status = None
    elif isinstance(raw, (int, float)) and math.isfinite(float(raw)):
        status = _CACHE_INTEGER_STATUSES.get(
            int(raw), "status_{0}".format(int(raw)))
    elif raw is not None:
        status = str(raw).strip().lower().replace("-", "_").replace(" ", "_")
        aliases = {
            "optimal": "ok",
            "loaded": "ok",
            "timelimit": "time_limit",
            "memory_limit": "oom",
            "out_of_memory": "oom",
            "mem_limit": "oom",
        }
        status = aliases.get(status, status)
    elapsed = result.get(time_field)
    try:
        elapsed = float(elapsed)
    except (TypeError, ValueError):
        elapsed = None
    if status is None and time_limit is not None and elapsed is not None:
        threshold = max(
            float(time_limit) - max(1.0, 0.001 * float(time_limit)), 0.0)
        if elapsed >= threshold:
            status = "time_limit"
    if status is None:
        status = "ok" if result.get("objective") is not None else "no_solution"
    return status


def is_cacheable_result(result):
    return isinstance(result, dict) and result.get("_status", "ok") != "error"

INSTANCE_DIR = "instances"
RESULT_FILE = "results_comparative_experiment.xlsx"
CACHE_FILE = "cache/cache_comparative_experiment.jsonl"
SCALE_ORDER = ["Small", "Medium", "Large"]

GUROBI_ENABLED = True
LSHAPED_ENABLED = True
BENDERS_ENABLED = True
HEURISTIC_ENABLED = True

GUROBI_MODULE = "Solvers.solver_gurobi"
LSHAPED_MODULE = "Solvers.solver_lshaped"
BENDERS_MODULE = "Solvers.solver_benders"
HEURISTIC_MODULE = "Solvers.solver_heuristic"

TIME_LIMIT = 3600.0
MIP_GAP = 0
THREADS = 0
HEURISTIC_THREADS = 4
SOLVER_OUTPUT = False
SOLVER_ERRORS_VISIBLE = False
SOFT_MEMORY_LIMIT_GB = None
NODEFILE_START_GB = None

BENDERS_MAX_ITERATIONS = 500
BENDERS_TOLERANCE = 1e-4
BENDERS_ROOT_ITERATIONS = 12
BENDERS_BLOCK_CACHE = 1024
BENDERS_EXACT_RECOURSE = True
BENDERS_SUBPROBLEM_GAP = 1e-6
BENDERS_ROOT_TIME_FRACTION = 0.5

HEURISTIC_TIME_LIMIT = 300.0
HEURISTIC_ROOT_ROUNDS = 10
HEURISTIC_INTEGER_ROUNDS = 6
HEURISTIC_LOCAL_ROUNDS = 2
HEURISTIC_MAX_NEIGHBORS = 12
HEURISTIC_SEARCH_SCENARIOS = 0
HEURISTIC_MASTER_GAP = 0.01
HEURISTIC_QUALITY_GAP = 0.005
HEURISTIC_ADAPTIVE = True

GUROBI_PARAMS = {
    "time_limit": TIME_LIMIT,
    "mip_gap": MIP_GAP,
    "threads": THREADS,
    "output": SOLVER_OUTPUT,
    "soft_memory_limit_gb": SOFT_MEMORY_LIMIT_GB,
    "nodefile_start_gb": NODEFILE_START_GB,
}

LSHAPED_PARAMS = {
    "time_limit": TIME_LIMIT,
    "mip_gap": MIP_GAP,
    "threads": THREADS,
    "output": SOLVER_OUTPUT,
}

BENDERS_PARAMS = {
    "time_limit": TIME_LIMIT,
    "mip_gap": MIP_GAP,
    "threads": THREADS,
    "output": SOLVER_OUTPUT,
    "max_iterations": BENDERS_MAX_ITERATIONS,
    "tolerance": BENDERS_TOLERANCE,
    "root_iterations": BENDERS_ROOT_ITERATIONS,
    "block_cache_size": BENDERS_BLOCK_CACHE,
    "exact_recourse": BENDERS_EXACT_RECOURSE,
    "subproblem_gap": BENDERS_SUBPROBLEM_GAP,
    "root_time_fraction": BENDERS_ROOT_TIME_FRACTION,
}

HEURISTIC_PARAMS = {
    "time_limit": HEURISTIC_TIME_LIMIT,
    "threads": HEURISTIC_THREADS,
    "output": SOLVER_OUTPUT,
    "root_rounds": HEURISTIC_ROOT_ROUNDS,
    "integer_rounds": HEURISTIC_INTEGER_ROUNDS,
    "local_rounds": HEURISTIC_LOCAL_ROUNDS,
    "max_neighbors": HEURISTIC_MAX_NEIGHBORS,
    "search_scenarios": HEURISTIC_SEARCH_SCENARIOS,
    "master_gap": HEURISTIC_MASTER_GAP,
    "quality_gap": HEURISTIC_QUALITY_GAP,
    "adaptive": HEURISTIC_ADAPTIVE,
    "lp_backend": "scipy",
}

YELLOW_FILL = PatternFill(
    start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid"
)


def parse_file_name(file_name):
    stem = os.path.splitext(file_name)[0]
    match = re.match(r"^(.*)_(\d+)$", stem)
    if match is None:
        return stem, None
    return match.group(1), int(match.group(2))


def scale_rank(scale):
    if scale in SCALE_ORDER:
        return SCALE_ORDER.index(scale)
    return len(SCALE_ORDER)


def collect_instance_files(directory):
    entries = []
    for file_name in os.listdir(directory):
        if not file_name.endswith(".json"):
            continue
        scale, suffix = parse_file_name(file_name)
        entries.append(
            (
                scale_rank(scale),
                scale,
                suffix if suffix is not None else 0,
                file_name,
            )
        )
    entries.sort(key=lambda item: (item[0], item[1], item[2], item[3]))
    ordered = []
    counters = {}
    for _, scale, _, file_name in entries:
        counters[scale] = counters.get(scale, 0) + 1
        ordered.append(
            (scale, counters[scale], os.path.join(directory, file_name))
        )
    return ordered


def build_output_layout():
    solver_specs = []
    if GUROBI_ENABLED:
        solver_specs.append(
            (
                "gurobi",
                GUROBI_MODULE,
                GUROBI_PARAMS,
                [
                    ("gurobi_Obj", "objective", 18, 2),
                    ("gurobi_Gap(%)", "gap", 14, 2),
                    ("gurobi_Time(s)", "time", 15, 2),
                ],
            )
        )
    if LSHAPED_ENABLED:
        solver_specs.append(
            (
                "lshaped",
                LSHAPED_MODULE,
                LSHAPED_PARAMS,
                [
                    ("lshaped_Obj", "objective", 18, 2),
                    ("lshaped_Gap(%)", "gap", 15, 2),
                    ("lshaped_Time(s)", "time", 16, 2),
                ],
            )
        )
    if BENDERS_ENABLED:
        solver_specs.append(
            (
                "benders",
                BENDERS_MODULE,
                BENDERS_PARAMS,
                [
                    ("benders_Obj", "objective", 18, 2),
                    ("benders_Gap(%)", "gap", 15, 2),
                    ("benders_Time(s)", "time", 16, 2),
                ],
            )
        )
    if HEURISTIC_ENABLED:
        solver_specs.append(
            (
                "heuristic",
                HEURISTIC_MODULE,
                HEURISTIC_PARAMS,
                [
                    ("heuristic_Obj", "objective", 18, 2),
                    ("heuristic_Time(s)", "time", 18, 2),
                ],
            )
        )

    columns = ["Scale", "Instance"]
    widths = [9, 12]
    for _, _, _, fields in solver_specs:
        columns.extend(field[0] for field in fields)
        widths.extend(field[2] for field in fields)
    yellow_columns = [
        index + 1
        for index, column in enumerate(columns)
        if column.endswith("_Time(s)")
    ]
    return solver_specs, columns, widths, yellow_columns


def create_workbook(path, columns, widths, yellow_columns):
    book = Workbook()
    sheet = book.active
    sheet.title = "results"
    sheet.append(columns)
    for column in yellow_columns:
        sheet.cell(row=1, column=column).fill = YELLOW_FILL
    for index, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(index + 1)].width = width + 2
    book.save(path)
    return book, sheet


def append_row(book, sheet, path, values, yellow_columns):
    values = list(values)
    for column in yellow_columns:
        values[column - 1] = cap_report_time(values[column - 1])
    sheet.append(values)
    row = sheet.max_row
    for column in yellow_columns:
        sheet.cell(row=row, column=column).fill = YELLOW_FILL
    book.save(path)


def render(value, digits):
    if value is None:
        return "-"
    return "{0:.{1}f}".format(value, digits)


def cap_report_time(value):
    if value is None:
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(value):
        return None
    return min(max(value, 0.0), TIME_LIMIT)


def json_scalar(value):
    if value is None:
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) else None


def slim_result(result, status="ok", error=None):
    result = result or {}
    return {
        "objective": json_scalar(result.get("objective")),
        "gap": json_scalar(result.get("gap")),
        "time": json_scalar(result.get("time")),
        "_status": status,
        "_error": error,
    }


def print_row(values, widths):
    cells = []
    for index, value in enumerate(values):
        text = str(value)
        if index == 0:
            cells.append(text.ljust(widths[index]))
        else:
            cells.append(text.rjust(widths[index]))
    print("  ".join(cells), flush=True)


def run_solver(module, instance, params):
    # Wrapper for solver execution with exception handling
    if module is None:
        return slim_result({}, status="disabled")
    started = time.perf_counter()
    try:
        raw = module.solve_instance(instance, params)
    except Exception as exc:
        status = exception_status(exc)
        if SOLVER_ERRORS_VISIBLE and status not in ("oom", "time_limit"):
            raise
        return slim_result(
            {"time": time.perf_counter() - started},
            status=status,
            error=error_text(exc),
        )
    status = normalized_solver_status(
        raw, time_limit=params.get("time_limit"), time_field="time")
    return slim_result(raw, status=status)


def main():
    cache = load_cache(CACHE_FILE)
    solver_specs, columns, widths, yellow_columns = build_output_layout()
    solvers = [
        (
            name,
            importlib.import_module(module_name),
            params,
            fields,
        )
        for name, module_name, params, fields in solver_specs
    ]
    files = collect_instance_files(INSTANCE_DIR)
    book, sheet = create_workbook(
        RESULT_FILE, columns, widths, yellow_columns
    )
    print("[cache] loaded {0} comparative units from {1}".format(
        len(cache), CACHE_FILE), flush=True)
    print_row(columns, widths)

    for scale, order, path in files:
        with open(path, "r", encoding="utf-8") as handle:
            instance = json.load(handle)

        label = "{0}_{1:02d}".format(scale, order)
        instance_name = instance.get("meta", {}).get(
            "name", os.path.splitext(os.path.basename(path))[0])

        results = {}
        for name, module, params, _ in solvers:
            key = cache_identity(name, instance_name)
            if key in cache:
                results[name] = cache[key]
                tag = "cache"
            else:
                results[name] = run_solver(module, instance, params)
                if is_cacheable_result(results[name]):
                    store_cache(CACHE_FILE, name, instance_name, results[name])
                    cache[key] = results[name]
                    tag = "terminal-cached" if results[name].get(
                        "_status") in ("time_limit", "oom") else "solved"
                else:
                    tag = "error-not-cached"
            print("[{0}] {1}: status={2} obj={3} ({4})".format(
                label, name, results[name].get("_status", "unknown"),
                results[name].get("objective"), tag), flush=True)
            gc.collect()

        del instance
        gc.collect()

        record = [scale, label]
        display = [scale, label]
        for name, _, _, fields in solvers:
            result = results[name]
            for _, result_key, _, digits in fields:
                value = result.get(result_key)
                if result_key == "time" and value is not None:
                    value = cap_report_time(value)
                record.append(value)
                display.append(render(value, digits))

        append_row(
            book, sheet, RESULT_FILE, record, yellow_columns
        )
        print_row(display, widths)

        del results
        gc.collect()


if __name__ == "__main__":
    main()
