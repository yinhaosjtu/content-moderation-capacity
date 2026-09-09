import gc
import json
import os

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from ComparativeExperiments.comparative_experiment import (
    cache_identity,
    cap_report_time,
    exception_status,
    is_cacheable_result,
    load_cache,
    normalized_solver_status,
    parse_file_name,
    render,
    slim_result,
    store_cache,
)


INSTANCE_DIR = "instances_scalability"
RESULT_FILE = "results_scalability_experiment.xlsx"
CACHE_FILE = "cache/cache_scalability_experiment.jsonl"

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

TIME_LIMIT = 3600.0
THREADS = 4
SOLVER_OUTPUT = False
SOLVER_ERRORS_VISIBLE = False

HEURISTIC_PARAMS = {
    "time_limit": TIME_LIMIT,
    "threads": THREADS,
    "output": SOLVER_OUTPUT,
    "adaptive": True,
    "lp_backend": "scipy",
}

SCALE_SEEDS = [1, 2, 3]

DIMENSION_FIELDS = [
    ("n_languages", "|L|"),
    ("n_categories", "|C|"),
    ("n_tiers", "|K|"),
    ("n_days", "|D|"),
    ("n_operating_points", "|P|"),
    ("n_scenarios", "|W|"),
]

YELLOW_FILL = PatternFill(
    start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid"
)


def collect_instance_files(directory):
    entries = []
    for file_name in os.listdir(directory):
        if not file_name.endswith(".json"):
            continue
        label, suffix = parse_file_name(file_name)
        entries.append((label, suffix if suffix is not None else 0, file_name))
    entries.sort(key=lambda item: (item[0], item[1], item[2]))
    return [(label, suffix, os.path.join(directory, file_name))
            for label, suffix, file_name in entries]


def create_workbook(path, columns, widths, yellow_columns):
    book = Workbook()
    sheet = book.active
    sheet.title = "results"
    sheet.append(columns)
    for column in yellow_columns:
        sheet.cell(row=1, column=column).fill = YELLOW_FILL
    for index, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(index + 1)].width = width + 2
    book.save(path)
    return book, sheet


def append_row(book, sheet, path, values, yellow_columns):
    values = list(values)
    for column in yellow_columns:
        values[column - 1] = cap_report_time(values[column - 1])
    sheet.append(values)
    row = sheet.max_row
    for column in yellow_columns:
        sheet.cell(row=row, column=column).fill = YELLOW_FILL
    book.save(path)


def print_row(values, widths):
    cells = []
    for index, value in enumerate(values):
        text = str(value)
        if index <= 1:
            cells.append(text.ljust(widths[index]))
        else:
            cells.append(text.rjust(widths[index]))
    print("  ".join(cells), flush=True)


WORKER_MODULE = "ComparativeExperiments._scalability_worker"

WORKER_TIMEOUT_GRACE = 600.0


def run_heuristic(instance_path, params):
    import subprocess
    import sys
    import time

    time_limit = float(params.get("time_limit", TIME_LIMIT))
    threads = int(params.get("threads", 0) or 0)
    command = [
        sys.executable, "-u", "-m", WORKER_MODULE,
        instance_path, repr(time_limit), str(threads),
    ]
    started = time.perf_counter()
    try:
        completed = subprocess.run(
            command,
            cwd=PROJECT_ROOT,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=time_limit + WORKER_TIMEOUT_GRACE,
        )
    except subprocess.TimeoutExpired:
        return slim_result(
            {"time": time.perf_counter() - started},
            status="time_limit",
            error="worker exceeded {0:.0f}s wall clock".format(
                time_limit + WORKER_TIMEOUT_GRACE),
        )

    elapsed = time.perf_counter() - started
    payload = None
    for line in reversed(completed.stdout.splitlines()):
        if line.startswith("__RESULT__"):
            try:
                payload = json.loads(line[len("__RESULT__"):])
            except (TypeError, ValueError):
                payload = None
            break

    if payload is None:
        status = "oom"
        detail = (completed.stderr or "").strip()
        if detail:
            status = exception_status(Exception(detail))
        return slim_result(
            {"time": elapsed},
            status=status,
            error="worker exit={0}; {1}".format(
                completed.returncode, detail[-500:] if detail else "no output"),
        )

    status = normalized_solver_status(
        payload, time_limit=time_limit, time_field="time")
    return slim_result(payload, status=status)


def main():
    cache = load_cache(CACHE_FILE)

    columns = ["Point"] + [label for _, label in DIMENSION_FIELDS] \
        + ["n_block", "Obj", "Time(s)", "Status"]
    widths = [8, 5, 5, 5, 5, 5, 6, 9, 18, 12, 12]
    yellow_columns = [columns.index("Time(s)") + 1]

    files = collect_instance_files(INSTANCE_DIR)
    book, sheet = create_workbook(RESULT_FILE, columns, widths, yellow_columns)
    print("[cache] loaded {0} scalability units from {1}".format(
        len(cache), CACHE_FILE), flush=True)
    print_row(columns, widths)

    point_summary = {}

    for label, _, path in files:
        with open(path, "r", encoding="utf-8") as handle:
            instance = json.load(handle)

        dim = instance.get("dimensions", {})
        dim_values = [int(dim.get(key, 0)) for key, _ in DIMENSION_FIELDS]
        n_block = int(dim.get("n_days", 0)) * int(dim.get("n_scenarios", 0))
        instance_name = instance.get("meta", {}).get(
            "name", os.path.splitext(os.path.basename(path))[0])

        key = cache_identity("heuristic", instance_name)
        if key in cache:
            result = cache[key]
            tag = "cache"
        else:
            result = run_heuristic(path, HEURISTIC_PARAMS)
            if is_cacheable_result(result):
                store_cache(CACHE_FILE, "heuristic", instance_name, result)
                cache[key] = result
                tag = "terminal-cached" if result.get(
                    "_status") in ("time_limit", "oom") else "solved"
            else:
                tag = "error-not-cached"

        status = result.get("_status", "unknown")
        objective = result.get("objective")
        solve_time = cap_report_time(result.get("time"))
        print("[{0}] status={1} obj={2} time={3} ({4})".format(
            instance_name, status, objective, solve_time, tag), flush=True)

        record = [label] + dim_values + [n_block, objective,
                                         solve_time, status]
        display = [label] + [str(v) for v in dim_values] + [
            str(n_block), render(objective, 2),
            render(solve_time, 2), status]
        append_row(book, sheet, RESULT_FILE, record, yellow_columns)
        print_row(display, widths)

        bucket = point_summary.setdefault(
            label, {"dims": dim_values, "n_block": n_block,
                    "times": [], "objs": []})
        if status == "ok" and solve_time is not None:
            bucket["times"].append(solve_time)
        if status == "ok" and objective is not None:
            bucket["objs"].append(objective)

        del instance, result
        gc.collect()

    sheet.append([])
    sheet.append(["Mean per scale point (ok seeds only)"])
    for label in sorted(point_summary):
        bucket = point_summary[label]
        mean_time = (sum(bucket["times"]) / len(bucket["times"])
                     if bucket["times"] else None)
        mean_obj = (sum(bucket["objs"]) / len(bucket["objs"])
                    if bucket["objs"] else None)
        record = [label] + bucket["dims"] + [
            bucket["n_block"], mean_obj,
            cap_report_time(mean_time),
            "{0}/{1} ok".format(len(bucket["times"]), len(SCALE_SEEDS))]
        append_row(book, sheet, RESULT_FILE, record, yellow_columns)

    print("[done] results written to {0}".format(RESULT_FILE), flush=True)


if __name__ == "__main__":
    main()
