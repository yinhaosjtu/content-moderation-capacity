
import json
import math
import os
import time

import numpy as np
import gurobipy as gp
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from Solvers import solver_benders as sb


_CACHE_INTEGER_STATUSES = {
    2: "ok",
    3: "infeasible",
    4: "inf_or_unbd",
    5: "unbounded",
    6: "cutoff",
    7: "iteration_limit",
    8: "node_limit",
    9: "time_limit",
    10: "solution_limit",
    11: "interrupted",
    12: "numeric",
    13: "suboptimal",
    15: "user_obj_limit",
    16: "work_limit",
    17: "oom",
}


def cache_identity(unit, instance):
    return str(unit), str(instance)


def load_cache(path):
    cache = {}
    if not os.path.exists(path):
        return cache
    with open(path, "r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                entry = json.loads(line)
            except (TypeError, ValueError):
                continue
            unit = entry.get("unit")
            instance = entry.get("instance")
            if isinstance(unit, str) and isinstance(instance, str) and "result" in entry:
                cache[cache_identity(unit, instance)] = entry["result"]
    return cache


def _cache_json_default(value):
    item = getattr(value, "item", None)
    if callable(item):
        return item()
    raise TypeError(repr(value))


def store_cache(path, unit, instance, result):
    cache_dir = os.path.dirname(path)
    if cache_dir and not os.path.exists(cache_dir):
        os.makedirs(cache_dir, exist_ok=True)

    payload = {"unit": str(unit), "instance": str(instance), "result": result}
    if os.path.exists(path) and os.path.getsize(path) > 0:
        with open(path, "rb+") as boundary:
            boundary.seek(-1, os.SEEK_END)
            if boundary.read(1) != b"\n":
                boundary.seek(0, os.SEEK_END)
                boundary.write(b"\n")
                boundary.flush()
                os.fsync(boundary.fileno())
    with open(path, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(
            payload,
            ensure_ascii=False,
            allow_nan=True,
            default=_cache_json_default,
        ) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def exception_status(error):
    text = "{0}: {1}".format(type(error).__name__, error).lower()
    oom_markers = (
        "out of memory", "out-of-memory", "memory limit", "mem_limit",
        "cannot allocate memory", "failed to allocate", "allocation failed",
        "gurobi error 10001",
    )
    if (isinstance(error, MemoryError)
            or getattr(error, "errno", None) == 10001
            or any(marker in text for marker in oom_markers)):
        return "oom"
    timeout_markers = ("time limit", "time-limit", "timed out", "timeout")
    if isinstance(error, TimeoutError) or any(
            marker in text for marker in timeout_markers):
        return "time_limit"
    return "error"


def error_text(error):
    return "{0}: {1}".format(type(error).__name__, error)


def normalized_solver_status(result, time_limit=None, time_field="time"):
    if not isinstance(result, dict):
        return "error"
    raw = result.get("status")
    status = None
    if isinstance(raw, bool):
        status = None
    elif isinstance(raw, (int, float)) and math.isfinite(float(raw)):
        status = _CACHE_INTEGER_STATUSES.get(
            int(raw), "status_{0}".format(int(raw)))
    elif raw is not None:
        status = str(raw).strip().lower().replace("-", "_").replace(" ", "_")
        aliases = {
            "optimal": "ok",
            "loaded": "ok",
            "timelimit": "time_limit",
            "memory_limit": "oom",
            "out_of_memory": "oom",
            "mem_limit": "oom",
        }
        status = aliases.get(status, status)
    elapsed = result.get(time_field)
    try:
        elapsed = float(elapsed)
    except (TypeError, ValueError):
        elapsed = None
    if status is None and time_limit is not None and elapsed is not None:
        threshold = max(
            float(time_limit) - max(1.0, 0.001 * float(time_limit)), 0.0)
        if elapsed >= threshold:
            status = "time_limit"
    if status is None:
        status = "ok" if result.get("objective") is not None else "no_solution"
    return status


def is_cacheable_result(result):
    return isinstance(result, dict) and result.get("_status", "ok") != "error"

INSTANCE_DIR = "instances"
SCALE = "Medium"


def discover_m1_path(directory, scale):
    best_ordinal = None
    best_path = None
    prefix = scale + "_"
    for name in os.listdir(directory):
        if not (name.startswith(prefix) and name.endswith(".json")):
            continue
        stem = os.path.splitext(name)[0]
        parts = stem.rsplit("_", 1)
        if len(parts) == 2 and parts[1].isdigit():
            ordinal = int(parts[1])
            if best_ordinal is None or ordinal < best_ordinal:
                best_ordinal = ordinal
                best_path = os.path.join(directory, name)
    return best_path


INSTANCE_PATH = discover_m1_path(INSTANCE_DIR, SCALE)
RESULT_FILE = "results_m1_anatomy.xlsx"
CACHE_FILE = "cache/cache_m1_anatomy.jsonl"

MIP_GAP = 1e-4
TIME_LIMIT = 3600.0
THREADS = 0

COST_KEYS = ("exp_penalty_cost", "exp_flex_labour_cost")
QUANTITY_KEYS = ("exp_outsourced_hours", "exp_overtime_hours",
                 "exp_inspection_hours", "exp_adjudicated_items",
                 "exp_unreviewed_release_items", "exp_unreviewed_removal_items",
                 "exp_shortfall_cell_items", "exp_shortfall_lng_net_items",
                 "exp_shortfall_qc_hours")


def solve():
    with open(INSTANCE_PATH, "r", encoding="utf-8") as handle:
        inst = json.load(handle)
    unit = "m1_anatomy"
    instance_name = inst["meta"]["name"]
    identity = cache_identity(unit, instance_name)
    cache = load_cache(CACHE_FILE)
    if identity in cache:
        return inst, cache[identity], "cache"
    started = time.perf_counter()
    try:
        res = sb.solve_instance(inst, {
            "mip_gap": MIP_GAP,
            "threads": THREADS,
            "time_limit": TIME_LIMIT,
            "extract_channels": True,
        })
        slim = {
            "objective": res.get("objective"),
            "headcount": res.get("headcount"),
            "disposition": res.get("disposition"),
            "channels": res.get("channels"),
            "time": res.get("time"),
            "_status": normalized_solver_status(res, TIME_LIMIT, "time"),
        }
    except Exception as error:
        status = exception_status(error)
        if status not in ("oom", "time_limit"):
            raise
        slim = {
            "objective": None,
            "headcount": None,
            "disposition": None,
            "channels": None,
            "time": time.perf_counter() - started,
            "_status": status,
            "_error": error_text(error),
        }
    if is_cacheable_result(slim):
        store_cache(CACHE_FILE, unit, instance_name, slim)
    return inst, slim, "solved"


def set_widths(sheet, headers):
    for col, head in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = max(
            12, min(len(str(head)) + 2, 28))


def write_headcount(book, inst, head):
    languages = inst["sets"]["languages"]
    tiers = inst["sets"]["tiers"]
    arr = np.asarray(head, dtype=int)
    sheet = book.active
    sheet.title = "headcount"
    headers = ["language"] + ["Tier {0}".format(t) for t in tiers] + ["total"]
    sheet.append(headers)
    set_widths(sheet, headers)
    for l, name in enumerate(languages):
        row = [int(arr[l, k]) for k in range(arr.shape[1])]
        sheet.append([name] + row + [int(sum(row))])
    col_tot = [int(arr[:, k].sum()) for k in range(arr.shape[1])]
    sheet.append(["total"] + col_tot + [int(arr.sum())])


def write_disposition(book, inst, dispo):
    languages = inst["sets"]["languages"]
    categories = inst["sets"]["categories"]
    sheet = book.create_sheet("disposition")
    headers = ["language"] + list(categories)
    sheet.append(headers)
    set_widths(sheet, headers)
    for l, name in enumerate(languages):
        sheet.append([name] + [int(v) for v in dispo[l]])


def write_cost_breakdown(book, inst, res):
    gamma = float(inst["uncertainty"]["gamma"])
    ch = res.get("channels") or {}
    objective = float(res["objective"])
    f_commit = float(ch.get("F_commit", 0.0))
    e_recourse = float(ch.get("E_recourse", 0.0))
    cvar95 = float(ch.get("cvar95", 0.0))

    committed = f_commit
    expected_recourse = (1.0 - gamma) * e_recourse
    tail = gamma * cvar95

    sheet = book.create_sheet("cost_breakdown")
    headers = ["component", "absolute", "share_of_objective_pct"]
    sheet.append(headers)
    set_widths(sheet, headers)

    def pct(value):
        return 100.0 * value / objective if objective else None

    sheet.append(["F_commit (first-stage)", committed, pct(committed)])
    sheet.append(["(1-gamma) * E_recourse", expected_recourse,
                  pct(expected_recourse)])
    sheet.append(["gamma * CVaR_0.95", tail, pct(tail)])
    sheet.append(["objective (sum)", objective, pct(objective)])
    sheet.append([])

    sheet.append(["recourse cost split", "absolute",
                  "share_of_E_recourse_pct"])
    for k in COST_KEYS:
        v = float(ch.get(k, 0.0))
        share = 100.0 * v / e_recourse if e_recourse else None
        sheet.append([k, v, share])
    sheet.append([])

    sheet.append(["expected quantity", "absolute", ""])
    for k in QUANTITY_KEYS:
        sheet.append([k, float(ch.get(k, 0.0)), ""])
    sheet.append([])

    sheet.append(["E_recourse (mean recourse)", e_recourse, ""])
    sheet.append(["CVaR_0.95 (recourse tail)", cvar95, ""])
    sheet.append(["gamma", gamma, ""])
    sheet.append(["mix_fraction", float(ch.get("mix_fraction", 0.0)), ""])


def main():
    inst, res, tag = solve()
    if res.get("objective") is None or res.get("headcount") is None:
        raise SystemExit("solve ended with status {0}: no objective/headcount returned".format(
            res.get("_status", "no_solution")))
    head = res["headcount"]
    dispo = res["disposition"]

    book = Workbook()
    write_headcount(book, inst, head)
    write_disposition(book, inst, dispo)
    write_cost_breakdown(book, inst, res)
    book.save(RESULT_FILE)

    arr = np.asarray(head, dtype=int)
    print("[M1 anatomy] instance={0} ({1}) objective={2}".format(
        inst["meta"]["name"], tag, res["objective"]), flush=True)
    print("[M1 anatomy] n_total={0} per_tier={1}".format(
        int(arr.sum()), [int(arr[:, k].sum()) for k in range(arr.shape[1])]),
        flush=True)
    print("done: {0}".format(RESULT_FILE), flush=True)


if __name__ == "__main__":
    main()
