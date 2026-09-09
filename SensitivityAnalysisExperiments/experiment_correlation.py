
import json
import math
import os
import time

import numpy as np
from numpy.polynomial.legendre import leggauss
from scipy.special import erf
from scipy.stats import gamma as gamma_dist
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from Solvers import solver_benders as sb


_CACHE_INTEGER_STATUSES = {
    2: "ok",
    3: "infeasible",
    4: "inf_or_unbd",
    5: "unbounded",
    6: "cutoff",
    7: "iteration_limit",
    8: "node_limit",
    9: "time_limit",
    10: "solution_limit",
    11: "interrupted",
    12: "numeric",
    13: "suboptimal",
    15: "user_obj_limit",
    16: "work_limit",
    17: "oom",
}


def cache_identity(unit, instance):
    return str(unit), str(instance)


def load_cache(path):
    cache = {}
    if not os.path.exists(path):
        return cache
    with open(path, "r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                entry = json.loads(line)
            except (TypeError, ValueError):
                continue
            unit = entry.get("unit")
            instance = entry.get("instance")
            if isinstance(unit, str) and isinstance(instance, str) and "result" in entry:
                cache[cache_identity(unit, instance)] = entry["result"]
    return cache


def _cache_json_default(value):
    item = getattr(value, "item", None)
    if callable(item):
        return item()
    raise TypeError(repr(value))


def store_cache(path, unit, instance, result):
    cache_dir = os.path.dirname(path)
    if cache_dir and not os.path.exists(cache_dir):
        os.makedirs(cache_dir, exist_ok=True)

    payload = {"unit": str(unit), "instance": str(instance), "result": result}
    if os.path.exists(path) and os.path.getsize(path) > 0:
        with open(path, "rb+") as boundary:
            boundary.seek(-1, os.SEEK_END)
            if boundary.read(1) != b"\n":
                boundary.seek(0, os.SEEK_END)
                boundary.write(b"\n")
                boundary.flush()
                os.fsync(boundary.fileno())
    with open(path, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(
            payload,
            ensure_ascii=False,
            allow_nan=True,
            default=_cache_json_default,
        ) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def exception_status(error):
    text = "{0}: {1}".format(type(error).__name__, error).lower()
    oom_markers = (
        "out of memory", "out-of-memory", "memory limit", "mem_limit",
        "cannot allocate memory", "failed to allocate", "allocation failed",
        "gurobi error 10001",
    )
    if (isinstance(error, MemoryError)
            or getattr(error, "errno", None) == 10001
            or any(marker in text for marker in oom_markers)):
        return "oom"
    timeout_markers = ("time limit", "time-limit", "timed out", "timeout")
    if isinstance(error, TimeoutError) or any(
            marker in text for marker in timeout_markers):
        return "time_limit"
    return "error"


def error_text(error):
    return "{0}: {1}".format(type(error).__name__, error)


def normalized_solver_status(result, time_limit=None, time_field="time"):
    if not isinstance(result, dict):
        return "error"
    raw = result.get("status")
    status = None
    if isinstance(raw, bool):
        status = None
    elif isinstance(raw, (int, float)) and math.isfinite(float(raw)):
        status = _CACHE_INTEGER_STATUSES.get(
            int(raw), "status_{0}".format(int(raw)))
    elif raw is not None:
        status = str(raw).strip().lower().replace("-", "_").replace(" ", "_")
        aliases = {
            "optimal": "ok",
            "loaded": "ok",
            "timelimit": "time_limit",
            "memory_limit": "oom",
            "out_of_memory": "oom",
            "mem_limit": "oom",
        }
        status = aliases.get(status, status)
    elapsed = result.get(time_field)
    try:
        elapsed = float(elapsed)
    except (TypeError, ValueError):
        elapsed = None
    if status is None and time_limit is not None and elapsed is not None:
        threshold = max(
            float(time_limit) - max(1.0, 0.001 * float(time_limit)), 0.0)
        if elapsed >= threshold:
            status = "time_limit"
    if status is None:
        status = "ok" if result.get("objective") is not None else "no_solution"
    return status


def is_cacheable_result(result):
    return isinstance(result, dict) and result.get("_status", "ok") != "error"

DEFAULT_MIP_GAP = 1e-4
DEFAULT_TIME_LIMIT = 3600.0


def solve_and_extract(instance, mip_gap=DEFAULT_MIP_GAP,
                      time_limit=DEFAULT_TIME_LIMIT, threads=0):
    res = sb.solve_instance(instance, {
        "mip_gap": mip_gap,
        "threads": threads,
        "time_limit": time_limit,
        "extract_channels": True,
    })
    out = {
        "objective": res.get("objective"),
        "bound": res.get("bound"),
        "gap_pct": res.get("gap"),
        "time_s": res.get("time"),
        "fractional_blocks": res.get("fractional_blocks"),
        "_status": normalized_solver_status(res, time_limit, "time"),
    }
    head = res.get("headcount")
    if head is not None:
        n_arr = np.asarray(head, dtype=int)
        out["n"] = head
        out["n_total"] = int(n_arr.sum())
        out["n_tier"] = [int(n_arr[:, k].sum()) for k in range(n_arr.shape[1])]
    out["y"] = res.get("disposition")
    ch = res.get("channels")
    if ch is not None:
        out.update(ch)
        out["total_expected"] = ch.get("F_commit", 0.0) + ch.get("E_recourse", 0.0)
    return out


CATEGORY_LIBRARY = [
    ("spam_and_platform_manipulation", 0.10, 0.05),
    ("adult_sexual_content", 0.30, 0.30),
    ("harassment_and_bullying", 0.45, 0.55),
    ("hate_speech", 0.62, 0.95),
    ("violence_and_incitement", 0.78, 0.70),
    ("terrorism_and_violent_extremism", 0.95, 0.60),
]

CATEGORY_SELECTION_ORDER = [0, 3, 5, 1, 4, 2]

GL_NODES, GL_WEIGHTS = leggauss(64)

DOW_PATTERN = np.array([1.00, 0.98, 0.99, 1.02, 1.07, 1.12, 1.06])

SCHEDULED_HOURS = 8.0

BODY_LEVELS = np.linspace(0.01, 0.90, 36)

TAIL_LEVELS = np.array([0.915, 0.930, 0.945, 0.955, 0.965, 0.973, 0.980,
                        0.986, 0.990, 0.993, 0.995, 0.9965, 0.998, 0.999,
                        0.9995, 0.9998, 0.99993, 0.99998, 0.999995,
                        0.9999995])

TAU0_LO_RANGE = (1.75, 1.05)
TAU0_HI_RANGE = (2.55, 3.25)
TAU1_LO_RANGE = (1.55, 0.85)
TAU1_HI_RANGE = (2.40, 2.95)


def std_normal_cdf(x):
    return 0.5 * (1.0 + erf(np.asarray(x, dtype=float) / math.sqrt(2.0)))


def bivariate_normal_cdf(h, k, r):
    h, k, r = np.broadcast_arrays(np.asarray(h, dtype=float),
                                  np.asarray(k, dtype=float),
                                  np.asarray(r, dtype=float))
    base = std_normal_cdf(h) * std_normal_cdf(k)
    half = 0.5 * r[..., None]
    node = half * (GL_NODES + 1.0)
    weight = half * GL_WEIGHTS
    comp = 1.0 - node * node
    hh = h[..., None]
    kk = k[..., None]
    dens = np.exp(-(hh * hh - 2.0 * node * hh * kk + kk * kk) / (2.0 * comp))
    dens = dens / (2.0 * math.pi * np.sqrt(comp))
    return base + np.sum(weight * dens, axis=-1)


def logit(x):
    return np.log(x / (1.0 - x))


def expit(x):
    return 1.0 / (1.0 + np.exp(-x))


def split_operating_grid(n_points):
    for n1 in (3, 2, 1):
        if n_points % n1 == 0 and n_points // n1 >= n1:
            return n_points // n1, n1
    return n_points, 1


def build_categories(n_categories, n_tiers):
    picks = sorted(CATEGORY_SELECTION_ORDER[:n_categories])
    names = [CATEGORY_LIBRARY[i][0] for i in picks]
    severity = np.array([CATEGORY_LIBRARY[i][1] for i in picks])
    expressive = np.array([CATEGORY_LIBRARY[i][2] for i in picks])
    k_min = np.minimum(n_tiers, 1 + np.floor(severity * n_tiers)).astype(int)
    k_min[int(np.argmin(severity))] = 1
    c_fc = [i for i in range(len(picks)) if severity[i] >= 0.90]
    if len(c_fc) == 0:
        c_fc = [int(np.argmax(severity))]
    c_fo = [i for i in range(len(picks))
            if abs(expressive[i] - 0.95) < 1e-9 and i not in c_fc]
    c_res = [i for i in range(len(picks)) if k_min[i] > 1]
    c_free = [i for i in range(len(picks)) if i not in c_fc and i not in c_fo]
    return names, severity, expressive, k_min, c_fc, c_fo, c_res, c_free


def build_operating_points(rng, n_points):
    n0, n1 = split_operating_grid(n_points)
    t0_lo = np.linspace(TAU0_LO_RANGE[0], TAU0_LO_RANGE[1], n0) + rng.uniform(-0.05, 0.05, n0)
    t0_hi = np.linspace(TAU0_HI_RANGE[0], TAU0_HI_RANGE[1], n0) + rng.uniform(-0.05, 0.05, n0)
    t1_lo = np.linspace(TAU1_LO_RANGE[0], TAU1_LO_RANGE[1], n1) + rng.uniform(-0.05, 0.05, n1)
    t1_hi = np.linspace(TAU1_HI_RANGE[0], TAU1_HI_RANGE[1], n1) + rng.uniform(-0.05, 0.05, n1)
    tau0_lo = np.zeros(n_points)
    tau0_hi = np.zeros(n_points)
    tau1_lo = np.zeros(n_points)
    tau1_hi = np.zeros(n_points)
    lvl0 = np.zeros(n_points, dtype=int)
    lvl1 = np.zeros(n_points, dtype=int)
    for i0 in range(n0):
        for i1 in range(n1):
            p = i0 * n1 + i1
            tau0_lo[p] = t0_lo[i0]
            tau0_hi[p] = t0_hi[i0]
            tau1_lo[p] = t1_lo[i1]
            tau1_hi[p] = t1_hi[i1]
            lvl0[p] = i0
            lvl1[p] = i1
    p_ref = (n0 // 2) * n1 + (n1 // 2)
    return tau0_lo, tau0_hi, tau1_lo, tau1_hi, lvl0, lvl1, int(p_ref), n0, n1


def outcome_probabilities(tau0_lo, tau0_hi, tau1_lo, tau1_hi, mu0, mu1, rho):
    x_lo = tau0_lo - mu0
    x_hi = tau0_hi - mu0
    y_lo = tau1_lo - mu1
    y_hi = tau1_hi - mu1
    phi_lo = std_normal_cdf(x_lo)
    phi_hi = std_normal_cdf(x_hi)
    hi_hi = bivariate_normal_cdf(x_hi, y_hi, rho)
    lo_hi = bivariate_normal_cdf(x_lo, y_hi, rho)
    hi_lo = bivariate_normal_cdf(x_hi, y_lo, rho)
    lo_lo = bivariate_normal_cdf(x_lo, y_lo, rho)
    p_b0 = 1.0 - phi_hi
    p_p0 = phi_lo
    p_esc = phi_hi - phi_lo
    p_mid = hi_hi - lo_hi - hi_lo + lo_lo
    p_p1 = hi_lo - lo_lo
    p_b1 = p_esc - p_mid - p_p1
    p_blk = p_b0 + p_b1
    p_pas = p_p0 + p_p1
    return p_esc, p_mid, p_blk, p_pas


def build_detection_coefficients(delta_grid, tau0_lo, tau0_hi, tau1_lo, tau1_hi,
                                 delta0, delta1, rho, kappa, pi0, bsens):
    grid = delta_grid[None, None, None, :]
    eps = 1.0 - np.exp(-kappa[:, :, None, None] * grid)
    mu0_one = (1.0 - eps) * delta0[:, :, None, None]
    mu1_one = (1.0 - eps) * delta1[:, :, None, None]
    zero = np.zeros_like(mu0_one)
    t0lo = tau0_lo[None, None, :, None]
    t0hi = tau0_hi[None, None, :, None]
    t1lo = tau1_lo[None, None, :, None]
    t1hi = tau1_hi[None, None, :, None]
    corr = rho[:, :, None, None]
    e0, m0, b0, s0 = outcome_probabilities(t0lo, t0hi, t1lo, t1hi, zero, zero, corr)
    e1, m1, b1, s1 = outcome_probabilities(t0lo, t0hi, t1lo, t1hi, mu0_one, mu1_one, corr)
    prev = expit(logit(pi0)[:, :, None, None] + bsens[:, :, None, None] * grid)
    phi_esc = prev * e1 + (1.0 - prev) * e0
    phi_mid = prev * m1 + (1.0 - prev) * m0
    phi_blk = prev * b1 + (1.0 - prev) * b0
    phi_pas = prev * s1 + (1.0 - prev) * s0
    e_fn = prev * s1
    e_fp = (1.0 - prev) * b0
    pi_hat = prev * m1 / phi_mid
    return phi_esc, phi_mid, phi_blk, phi_pas, e_fn, e_fp, pi_hat, prev


def build_scenarios(rng, n_days, n_scen, sigma_z, kappa_delta, beta_delta,
                    rho_zd, varrho, delta_grid, grid_levels):
    cov = np.array([[1.0, rho_zd], [rho_zd, 1.0]])
    chol = np.linalg.cholesky(cov)
    state = rng.standard_normal((n_scen, 2)) @ chol.T
    scores = np.zeros((n_days, n_scen, 2))
    for d in range(n_days):
        innov = rng.standard_normal((n_scen, 2)) @ chol.T
        state = varrho * state + math.sqrt(1.0 - varrho * varrho) * innov
        scores[d] = state
    z_val = np.exp(-0.5 * sigma_z * sigma_z + sigma_z * scores[:, :, 0])
    u_delta = std_normal_cdf(scores[:, :, 1])
    u_delta = np.clip(u_delta, 1e-12, 1.0 - 1e-12)
    pos = np.searchsorted(grid_levels, u_delta)
    pos = np.clip(pos, 1, len(grid_levels) - 1)
    left = grid_levels[pos - 1]
    right = grid_levels[pos]
    idx = np.where(u_delta - left <= right - u_delta, pos - 1, pos)
    return z_val, idx.astype(int), delta_grid[idx]


def to_list(arr, digits):
    return np.round(np.asarray(arr, dtype=float), digits).tolist()


def build_instance(seed, size, label):
    rng = np.random.default_rng(int(seed))
    n_l = int(size["n_languages"])
    n_c = int(size["n_categories"])
    n_k = int(size["n_tiers"])
    n_d = int(size["n_days"])
    n_p = int(size["n_operating_points"])
    n_w = int(size["n_scenarios"])

    names_c, severity, expressive, k_min, c_fc, c_fo, c_res, c_free = build_categories(n_c, n_k)

    pool_max = 8.0e6
    pool_min = 1.2e5
    pool_base = np.exp(np.linspace(math.log(pool_max), math.log(pool_min), n_l))
    pool_base = pool_base * np.exp(rng.normal(0.0, 0.12, n_l))
    pool_base = np.sort(pool_base)[::-1]
    span = max(math.log(pool_base.max()) - math.log(pool_base.min()), 1e-9)
    resource = (np.log(pool_base) - math.log(pool_base.min())) / span

    amp = rng.uniform(0.03, 0.10, n_l)
    phase = rng.uniform(0.0, 2.0 * math.pi, n_l)
    day_index = np.arange(n_d)
    dow = DOW_PATTERN / DOW_PATTERN.mean()
    seasonal = 1.0 + amp[:, None] * np.sin(2.0 * math.pi * day_index[None, :] / max(n_d, 1) + phase[:, None])
    a_bar = pool_base[:, None] * dow[day_index % 7][None, :] * seasonal

    scope_total = rng.uniform(0.25, 0.55, n_l)
    cat_weight = 1.0 / (1.0 + 2.2 * severity)
    cat_weight = cat_weight / cat_weight.sum()
    psi = np.zeros((n_l, n_c))
    for l in range(n_l):
        psi[l] = scope_total[l] * rng.dirichlet(4.0 * n_c * cat_weight)

    pi_center = np.exp(math.log(0.060) + (math.log(0.004) - math.log(0.060)) * severity)
    pi0 = np.clip(pi_center[None, :] * np.exp(rng.normal(0.0, 0.25, (n_l, n_c))), 0.001, 0.150)

    delta0 = (1.75 + 0.75 * resource)[:, None] * (1.0 - 0.15 * severity)[None, :]
    delta0 = delta0 * np.exp(rng.normal(0.0, 0.06, (n_l, n_c)))
    delta1 = np.minimum(delta0 + rng.uniform(0.50, 1.10, (n_l, n_c)), 3.50)
    rho_layer = rng.uniform(0.35, 0.75, (n_l, n_c))
    kappa = (0.62 - 0.28 * resource)[:, None] * (1.0 + 0.35 * severity)[None, :]
    kappa = np.clip(kappa * np.exp(rng.normal(0.0, 0.12, (n_l, n_c))), 0.10, 1.20)
    bsens = rng.uniform(0.25, 0.95, (n_l, n_c)) * (1.0 + 0.40 * severity)[None, :]

    tau0_lo, tau0_hi, tau1_lo, tau1_hi, lvl0, lvl1, p_ref, n0, n1 = build_operating_points(rng, n_p)

    sigma_z = float(rng.uniform(0.18, 0.35))
    kappa_delta = float(0.5 * rng.uniform(0.45, 0.80))
    beta_delta = float(2.0 * rng.uniform(0.30, 0.60))
    rho_zd_draw = float(rng.uniform(0.50, 0.75))
    rho_zd = rho_zd_draw if size.get("rho_z_delta") is None else float(size["rho_z_delta"])
    varrho = float(rng.uniform(0.55, 0.80))
    zeta1_draw = float(rng.uniform(0.30, 0.90))
    zeta1 = zeta1_draw if size.get("zeta1") is None else float(size["zeta1"])
    zeta2_draw = float(rng.uniform(0.80, 1.60))
    zeta2 = zeta2_draw if size.get("zeta2") is None else float(size["zeta2"])

    grid_levels = np.concatenate([[0.0], BODY_LEVELS, TAIL_LEVELS])
    delta_grid = gamma_dist.ppf(grid_levels, a=kappa_delta, scale=beta_delta)
    delta_grid[0] = 0.0
    n_g = len(delta_grid)

    phi_esc, phi_mid, phi_blk, phi_pas, e_fn, e_fp, pi_hat, prev = build_detection_coefficients(
        delta_grid, tau0_lo, tau0_hi, tau1_lo, tau1_hi, delta0, delta1, rho_layer, kappa, pi0, bsens)

    z_val, delta_idx, delta_val = build_scenarios(
        rng, n_d, n_w, sigma_z, kappa_delta, beta_delta, rho_zd, varrho,
        delta_grid, grid_levels)

    theta_cat = (20.0 + 140.0 * severity) / 3600.0
    theta_lang = rng.uniform(0.90, 1.15, n_l)
    tier_speed = 1.0 + 0.14 * np.arange(n_k)
    theta = theta_cat[None, :, None] * theta_lang[:, None, None] * tier_speed[None, None, :]
    theta = theta * np.exp(rng.normal(0.0, 0.08, (n_l, n_c, n_k)))

    chi_top = rng.uniform(0.74, 0.84, n_l)
    chi_in = chi_top[:, None] * (1.0 - 0.03 * np.arange(n_k)[None, :])
    chi_out = np.maximum(chi_top - rng.uniform(0.05, 0.11, n_l), 0.50)

    theta_in = theta / chi_in[:, None, :]
    theta_out = theta[:, :, 0] / chi_out[:, None]
    theta_qc = theta[:, :, n_k - 1] * rng.uniform(0.35, 0.60, (n_l, n_c)) / chi_in[:, n_k - 1][:, None]

    varsigma_blk = np.clip(rng.uniform(0.020, 0.050, n_c) * (1.0 + 0.5 * severity), 0.010, 0.100)
    varsigma_pas = np.clip(rng.uniform(0.0005, 0.0030, n_c) * (1.0 + 0.8 * severity), 0.0002, 0.0080)
    varsigma_pas = np.minimum(varsigma_pas, varsigma_blk)
    varsigma_adj_in = rng.uniform(0.020, 0.050, n_c)
    varsigma_adj_out = np.minimum(varsigma_adj_in + rng.uniform(0.020, 0.050, n_c), 0.150)

    o_bar_l = rng.uniform(0.6, 1.2, n_l)
    o_bar = np.repeat(o_bar_l[:, None], n_k, axis=1)

    acc_in1 = 0.905 - 0.075 * severity[None, :] + 0.045 * resource[:, None]
    acc_in1 = np.clip(acc_in1 + rng.normal(0.0, 0.012, (n_l, n_c)), 0.600, 0.960)
    step = rng.uniform(0.015, 0.035, (n_l, n_c, n_k))
    step[:, :, 0] = 0.0
    acc_in = np.clip(acc_in1[:, :, None] + np.cumsum(step, axis=2), 0.550, 0.985)
    acc_out = np.clip(acc_in[:, :, 0] - rng.uniform(0.020, 0.050, (n_l, n_c)), 0.550, 0.980)

    wage1 = rng.uniform(13.0, 32.0, n_l)
    tier_wage = 1.0 + 0.35 * np.arange(n_k)
    wage = wage1[:, None] * tier_wage[None, :]
    c_hr = wage * SCHEDULED_HOURS * n_d
    c_ot = 1.5 * wage
    c_out = wage1 * rng.uniform(1.40, 1.90, n_l)
    c_ml = float(rng.uniform(8.0e-4, 3.0e-3))
    c_fn = np.exp(math.log(1.5) + (math.log(200.0) - math.log(1.5)) * severity)
    c_fn = c_fn * np.exp(rng.normal(0.0, 0.18, n_c))
    c_fp = np.exp(math.log(1.5) + (math.log(22.0) - math.log(1.5)) * expressive)
    c_fp = c_fp * np.exp(rng.normal(0.0, 0.18, n_c))

    phi_mid_ref0 = phi_mid[:, :, p_ref, 0]
    frontline = np.array([1.0 if k_min[c] == 1 else 0.0 for c in range(n_c)])
    nominal_out_hours = np.einsum('lc,ld,lc,c->ld', psi, a_bar, phi_mid_ref0 * theta_out, frontline)
    depth = 0.25 * np.clip(0.03 + 0.55 * resource, 0.0, 1.0) * rng.uniform(0.80, 1.20, n_l)
    v0_bar = depth[:, None] * nominal_out_hours

    v_bar = v0_bar[:, :, None] / (1.0 + zeta1 * np.maximum(z_val - 1.0, 0.0)[None, :, :]
                                  + zeta2 * delta_val[None, :, :])

    phi_mid_best = phi_mid.max(axis=2)
    mid_by_scen = phi_mid_best[:, :, delta_idx]
    vol = a_bar[:, None, :, None] * z_val[None, None, :, :]
    u_bar = psi[:, :, None] * np.max(vol * mid_by_scen[:, :, :, :], axis=3)

    nominal_hours = np.zeros((n_l, n_k))
    for k in range(n_k):
        mask = np.array([1.0 if k_min[c] == k + 1 else 0.0 for c in range(n_c)])
        hrs = np.einsum('lc,ld,lc,c->ld', psi, a_bar, phi_mid_ref0 * theta_in[:, :, k], mask)
        nominal_hours[:, k] = hrs.max(axis=1)
    qc_hours = np.einsum('lc,ld,lc->ld', psi, a_bar,
                         phi_mid_ref0 * theta_qc * (varsigma_adj_in[None, :] + 0.0)) \
        + np.einsum('lc,ld,lc->ld', psi, a_bar, theta_qc * varsigma_blk[None, :] * phi_blk[:, :, p_ref, 0]) \
        + np.einsum('lc,ld,lc->ld', psi, a_bar, theta_qc * varsigma_pas[None, :] * phi_pas[:, :, p_ref, 0])
    nominal_hours[:, n_k - 1] = nominal_hours[:, n_k - 1] + qc_hours.max(axis=1)
    n_bar = np.maximum(np.ceil(3.0 * nominal_hours / SCHEDULED_HOURS), 5.0).astype(int)
    b_bar = float((c_hr * n_bar).sum())

    epsilon_lng = float(rng.uniform(0.03, 0.10))
    eta_cell = np.zeros((n_l, n_c))
    for c in c_res:
        eta_cell[:, c] = np.clip(0.68 + 0.30 * severity[c] + rng.uniform(-0.05, 0.05, n_l), 0.40, 0.95)
    varpi_lng = float(rng.uniform(10.0, 25.0))
    varpi_cell = varpi_lng * (1.4 + 1.8 * severity) * rng.uniform(0.95, 1.15, n_c)
    varpi_cell = np.maximum(varpi_cell, varpi_lng)
    qc_unit_penalty = float(rng.uniform(15.0, 45.0))
    varpi_qc = float(qc_unit_penalty / float(theta_qc.mean()))

    alpha = float(size.get("alpha", 0.95))
    gamma_risk = float(size.get("gamma", 0.30))

    instance = {
        "meta": {
            "name": "{0}_{1:02d}".format(label, int(seed)),
            "scale": label,
            "seed": int(seed),
            "model": "two_stage_risk_averse_content_moderation",
            "currency": "USD",
            "time_unit": "hour",
            "volume_unit": "item",
        },
        "dimensions": {
            "n_languages": n_l,
            "n_categories": n_c,
            "n_tiers": n_k,
            "n_days": n_d,
            "n_operating_points": n_p,
            "n_scenarios": n_w,
            "n_delta_grid": n_g,
        },
        "sets": {
            "languages": ["L{0:02d}".format(i + 1) for i in range(n_l)],
            "categories": names_c,
            "tiers": [k + 1 for k in range(n_k)],
            "channels": ["in", "out"],
            "days": [d + 1 for d in range(n_d)],
            "operating_points": [p for p in range(n_p)],
            "scenarios": [w for w in range(n_w)],
            "available_channel_tier": [["in", k + 1] for k in range(n_k)] + [["out", 1]],
            "k_min": [int(v) for v in k_min],
            "C_k": {str(k + 1): [c for c in range(n_c) if k_min[c] <= k + 1] for k in range(n_k)},
            "A_c": {str(c): [["in", k + 1] for k in range(n_k) if k + 1 >= k_min[c]]
                    + ([["out", 1]] if k_min[c] == 1 else []) for c in range(n_c)},
            "C_res": [int(c) for c in c_res],
            "C_fc": [int(c) for c in c_fc],
            "C_fo": [int(c) for c in c_fo],
            "C_free": [int(c) for c in c_free],
            "category_severity": to_list(severity, 4),
            "category_expressiveness": to_list(expressive, 4),
            "language_resource_level": to_list(resource, 6),
        },
        "operating_points": {
            "tau0_lo": to_list(tau0_lo, 6),
            "tau0_hi": to_list(tau0_hi, 6),
            "tau1_lo": to_list(tau1_lo, 6),
            "tau1_hi": to_list(tau1_hi, 6),
            "l0_level": [int(v) for v in lvl0],
            "l1_level": [int(v) for v in lvl1],
            "n_l0_levels": int(n0),
            "n_l1_levels": int(n1),
            "p_ref": int(p_ref),
        },
        "latent_score_model": {
            "Delta0": to_list(delta0, 6),
            "Delta1": to_list(delta1, 6),
            "rho": to_list(rho_layer, 6),
            "kappa": to_list(kappa, 6),
            "pi0": to_list(pi0, 8),
            "b": to_list(bsens, 6),
        },
        "delta_grid": to_list(delta_grid, 10),
        "delta_grid_levels": to_list(grid_levels, 10),
        "detection_coefficients": {
            "phi_esc": to_list(phi_esc, 10),
            "phi_mid": to_list(phi_mid, 10),
            "phi_blk": to_list(phi_blk, 10),
            "phi_pas": to_list(phi_pas, 10),
            "e_fn": to_list(e_fn, 10),
            "e_fp": to_list(e_fp, 10),
            "pi_hat": to_list(pi_hat, 10),
        },
        "volume": {
            "a_bar": to_list(a_bar, 2),
            "psi": to_list(psi, 8),
        },
        "workload": {
            "H": SCHEDULED_HOURS,
            "theta": to_list(theta, 8),
            "chi_in": to_list(chi_in, 6),
            "chi_out": to_list(chi_out, 6),
            "theta_in": to_list(theta_in, 8),
            "theta_out": to_list(theta_out, 8),
            "theta_qc": to_list(theta_qc, 8),
            "o_bar": to_list(o_bar, 4),
            "v0_bar": to_list(v0_bar, 4),
            "accuracy_in": to_list(acc_in, 6),
            "accuracy_out": to_list(acc_out, 6),
        },
        "inspection": {
            "varsigma_blk": to_list(varsigma_blk, 6),
            "varsigma_pas": to_list(varsigma_pas, 6),
            "varsigma_adj_in": to_list(varsigma_adj_in, 6),
            "varsigma_adj_out": to_list(varsigma_adj_out, 6),
        },
        "compliance": {
            "epsilon_lng": epsilon_lng,
            "eta_cell": to_list(eta_cell, 6),
            "varpi_lng": varpi_lng,
            "varpi_cell": to_list(varpi_cell, 4),
            "varpi_qc": round(varpi_qc, 4),
            "qc_unit_penalty": qc_unit_penalty,
        },
        "cost": {
            "c_hr": to_list(c_hr, 2),
            "c_ot": to_list(c_ot, 4),
            "c_out": to_list(c_out, 4),
            "c_ml": c_ml,
            "c_fn": to_list(c_fn, 4),
            "c_fp": to_list(c_fp, 4),
            "wage_straight_time": to_list(wage, 4),
            "n_bar": [[int(v) for v in row] for row in n_bar],
            "B_bar": round(b_bar, 2),
        },
        "uncertainty": {
            "sigma_Z": sigma_z,
            "kappa_delta": kappa_delta,
            "beta_delta": beta_delta,
            "rho_z_delta": rho_zd,
            "varrho": varrho,
            "zeta1": zeta1,
            "zeta2": zeta2,
            "alpha": alpha,
            "gamma": gamma_risk,
        },
        "scenarios": {
            "prob": [1.0 / n_w for _ in range(n_w)],
            "Z": to_list(z_val, 8),
            "delta": to_list(delta_val, 8),
            "delta_index": [[int(v) for v in row] for row in delta_idx],
            "v_bar": to_list(v_bar, 4),
        },
        "derived": {
            "U_bar": to_list(u_bar, 4),
        },
    }
    return instance


def validate_instance(inst):
    dim = inst["dimensions"]
    n_l = dim["n_languages"]
    n_c = dim["n_categories"]
    n_k = dim["n_tiers"]
    n_p = dim["n_operating_points"]
    psi = np.array(inst["volume"]["psi"])
    assert np.all(psi > 0.0)
    assert np.all(psi.sum(axis=1) <= 1.0 + 1e-9)
    phi_esc = np.array(inst["detection_coefficients"]["phi_esc"])
    phi_mid = np.array(inst["detection_coefficients"]["phi_mid"])
    phi_blk = np.array(inst["detection_coefficients"]["phi_blk"])
    phi_pas = np.array(inst["detection_coefficients"]["phi_pas"])
    e_fn = np.array(inst["detection_coefficients"]["e_fn"])
    e_fp = np.array(inst["detection_coefficients"]["e_fp"])
    pi_hat = np.array(inst["detection_coefficients"]["pi_hat"])
    assert np.all(phi_mid > 1e-9)
    assert np.all(phi_esc >= -1e-9) and np.all(phi_esc <= 1.0 + 1e-9)
    assert np.max(np.abs(phi_blk + phi_pas + phi_mid - 1.0)) < 1e-6
    assert np.all(phi_mid <= phi_esc + 1e-9)
    assert np.all(e_fn >= -1e-9) and np.all(e_fp >= -1e-9)
    assert np.all(pi_hat > 0.0) and np.all(pi_hat < 1.0)
    tau0_lo = np.array(inst["operating_points"]["tau0_lo"])
    tau0_hi = np.array(inst["operating_points"]["tau0_hi"])
    tau1_lo = np.array(inst["operating_points"]["tau1_lo"])
    tau1_hi = np.array(inst["operating_points"]["tau1_hi"])
    assert np.all(tau0_lo <= tau0_hi) and np.all(tau1_lo <= tau1_hi)
    assert 0 <= inst["operating_points"]["p_ref"] < n_p
    chi_in = np.array(inst["workload"]["chi_in"])
    chi_out = np.array(inst["workload"]["chi_out"])
    assert np.all(chi_in[:, 0] >= chi_out)
    assert np.all(chi_in > 0.0) and np.all(chi_in < 1.0)
    acc_in = np.array(inst["workload"]["accuracy_in"])
    acc_out = np.array(inst["workload"]["accuracy_out"])
    assert np.all(np.diff(acc_in, axis=2) >= -1e-12)
    assert np.all(acc_in[:, :, 0] >= acc_out)
    assert np.all(acc_out > 0.5)
    v_blk = np.array(inst["inspection"]["varsigma_blk"])
    v_pas = np.array(inst["inspection"]["varsigma_pas"])
    v_ai = np.array(inst["inspection"]["varsigma_adj_in"])
    v_ao = np.array(inst["inspection"]["varsigma_adj_out"])
    assert np.all(v_blk >= v_pas) and np.all(v_ao >= v_ai)
    c_hr = np.array(inst["cost"]["c_hr"])
    assert np.all(np.diff(c_hr, axis=1) > 0.0)
    varpi_cell = np.array(inst["compliance"]["varpi_cell"])
    assert np.all(varpi_cell >= inst["compliance"]["varpi_lng"] - 1e-9)
    assert 0.0 <= inst["compliance"]["epsilon_lng"] < 1.0
    unc = inst["uncertainty"]
    assert abs(unc["varrho"]) < 1.0
    assert 0.0 < unc["alpha"] < 1.0
    assert 0.0 <= unc["gamma"] < 1.0
    assert unc["zeta1"] >= 0.0 and unc["zeta2"] >= 0.0
    assert set(inst["sets"]["C_fc"]).isdisjoint(set(inst["sets"]["C_fo"]))
    assert len(inst["sets"]["C_k"][str(1)]) >= 1
    assert sorted(inst["sets"]["C_k"][str(n_k)]) == list(range(n_c))
    for c in range(n_c):
        for pair in inst["sets"]["A_c"][str(c)]:
            assert pair[0] in ("in", "out")
            if pair[0] == "out":
                assert pair[1] == 1
            assert pair[1] >= inst["sets"]["k_min"][c]
    z_val = np.array(inst["scenarios"]["Z"])
    delta_val = np.array(inst["scenarios"]["delta"])
    assert np.all(z_val > 0.0) and np.all(delta_val >= 0.0)
    v_bar = np.array(inst["scenarios"]["v_bar"])
    v0_bar = np.array(inst["workload"]["v0_bar"])
    assert np.all(v_bar <= v0_bar[:, :, None] + 1e-6)
    assert abs(sum(inst["scenarios"]["prob"]) - 1.0) < 1e-9
    u_bar = np.array(inst["derived"]["U_bar"])
    assert np.all(u_bar > 0.0)
    assert len(inst["delta_grid"]) == dim["n_delta_grid"]
    assert phi_mid.shape == (n_l, n_c, n_p, dim["n_delta_grid"])
    return True


MEDIUM_SIZE = {
    "n_languages": 5,
    "n_categories": 4,
    "n_tiers": 3,
    "n_days": 10,
    "n_operating_points": 6,
    "n_scenarios": 400,
    "alpha": 0.95,
    "gamma": 0.30,
    "rho_z_delta": None,
    "zeta1": None,
    "zeta2": None,
}


SCALE_LABEL = "Medium"
SCALE_SIZE = MEDIUM_SIZE
INSTANCE_DIR = "instances"
N_SEEDS = 5


def discover_ordinals(directory, scale, n_ordinals):
    ordinals = []
    prefix = scale + "_"
    for name in os.listdir(directory):
        if not (name.startswith(prefix) and name.endswith(".json")):
            continue
        stem = os.path.splitext(name)[0]
        parts = stem.rsplit("_", 1)
        if len(parts) == 2 and parts[1].isdigit():
            ordinals.append(int(parts[1]))
    ordinals = sorted(set(ordinals))
    return ordinals[:n_ordinals]


ORDINALS = discover_ordinals(INSTANCE_DIR, SCALE_LABEL, N_SEEDS)

RESULT_FILE = "results_correlation.xlsx"
CACHE_FILE = "cache/cache_correlation.jsonl"

RHO_VALUES = [0.0, 0.25, 0.50, 0.75]

BASE_GAMMA = 0.30
BASE_ALPHA = 0.95
MIP_GAP = 1e-4
TIME_LIMIT = 3600.0


def build_config_instance(seed, overrides):
    size = dict(SCALE_SIZE)
    size["gamma"] = BASE_GAMMA
    size["alpha"] = BASE_ALPHA
    size.update(overrides)
    inst = build_instance(seed, size, SCALE_LABEL)
    validate_instance(inst)
    return inst


def metrics(result):
    n_tier = result.get("n_tier") or []
    n_total = result.get("n_total")
    senior_share = (100.0 * n_tier[-1] / n_total
                    if n_tier and n_total else None)
    return {
        "objective": result.get("objective"),
        "n_total": n_total,
        "n_senior": n_tier[-1] if n_tier else None,
        "senior_share_pct": senior_share,
        "cvar95": result.get("cvar95"),
        "exp_penalty": result.get("exp_penalty_cost"),
        "exp_outsourced_hours": result.get("exp_outsourced_hours"),
        "exp_unreviewed_release": result.get("exp_unreviewed_release_items"),
    }


def run_config(cache, unit, instance_name, seed, overrides):
    identity = cache_identity(unit, instance_name)
    if identity in cache:
        return cache[identity]
    started = time.perf_counter()
    try:
        inst = build_config_instance(seed, overrides)
        result = solve_and_extract(inst, mip_gap=MIP_GAP, time_limit=TIME_LIMIT)
    except Exception as error:
        status = exception_status(error)
        if status not in ("oom", "time_limit"):
            raise
        result = {
            "objective": None,
            "time_s": time.perf_counter() - started,
            "_status": status,
            "_error": error_text(error),
        }
    if is_cacheable_result(result):
        store_cache(CACHE_FILE, unit, instance_name, result)
        cache[identity] = result
    return result


RAW_HEADERS = ["instance", "sweep", "config",
               "objective", "n_total", "n_senior", "senior_share_pct",
               "cvar95", "exp_penalty", "exp_outsourced_hours",
               "exp_unreviewed_release"]

INDEX_HEADERS = ["instance", "sweep", "config",
                 "n_total_idx", "n_senior_idx", "senior_share_pct",
                 "cvar95_idx", "penalty_idx", "outsourced_idx",
                 "release_idx"]


def index_row(m, base):
    def ratio(key):
        b = base.get(key)
        v = m.get(key)
        if b in (None, 0.0) or v is None:
            return None
        return 100.0 * v / b
    return {
        "n_total_idx": ratio("n_total"),
        "n_senior_idx": ratio("n_senior"),
        "senior_share_pct": m.get("senior_share_pct"),
        "cvar95_idx": ratio("cvar95"),
        "penalty_idx": ratio("exp_penalty"),
        "outsourced_idx": ratio("exp_outsourced_hours"),
        "release_idx": ratio("exp_unreviewed_release"),
    }


def main():
    cache = load_cache(CACHE_FILE)
    book = Workbook()
    raw = book.active
    raw.title = "raw"
    raw.append(RAW_HEADERS)
    idx = book.create_sheet("indexed")
    idx.append(INDEX_HEADERS)
    for sheet, heads in ((raw, RAW_HEADERS), (idx, INDEX_HEADERS)):
        for col, head in enumerate(heads, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = max(
                10, min(len(head) + 2, 20))
    book.save(RESULT_FILE)

    for ordinal_idx, ordinal in enumerate(ORDINALS, start=1):
        label = "M{0}".format(ordinal_idx)
        instance_name = "{0}_{1:02d}".format(SCALE_LABEL, ordinal)

        rho_metrics = {}
        for rho in RHO_VALUES:
            unit = "rho={0:.4f}".format(rho)
            result = run_config(
                cache, unit, instance_name, ordinal, {"rho_z_delta": rho})
            m = metrics(result)
            rho_metrics[rho] = m
            raw.append([label, "rho", "rho={0:.2f}".format(rho)]
                       + [m[k] for k in ("objective", "n_total", "n_senior",
                                         "senior_share_pct", "cvar95",
                                         "exp_penalty", "exp_outsourced_hours",
                                         "exp_unreviewed_release")])
            print("[{0}] rho={1:.2f} obj={2}".format(
                label, rho, m["objective"]), flush=True)
        base_rho = rho_metrics[0.0]
        for rho in RHO_VALUES:
            ir = index_row(rho_metrics[rho], base_rho)
            idx.append([label, "rho", "rho={0:.2f}".format(rho)]
                       + [ir[k] for k in ("n_total_idx", "n_senior_idx",
                                          "senior_share_pct", "cvar95_idx",
                                          "penalty_idx", "outsourced_idx",
                                          "release_idx")])

        zeta_metrics = {}
        for tag, ov in (("fixed", {"zeta1": 0.0, "zeta2": 0.0}),
                        ("shrink", {})):
            unit = "zeta={0}".format(tag)
            result = run_config(cache, unit, instance_name, ordinal, ov)
            m = metrics(result)
            zeta_metrics[tag] = m
            raw.append([label, "zeta", tag]
                       + [m[k] for k in ("objective", "n_total", "n_senior",
                                         "senior_share_pct", "cvar95",
                                         "exp_penalty", "exp_outsourced_hours",
                                         "exp_unreviewed_release")])
            print("[{0}] zeta={1} obj={2}".format(label, tag, m["objective"]),
                  flush=True)
        base_zeta = zeta_metrics["fixed"]
        for tag in ("fixed", "shrink"):
            ir = index_row(zeta_metrics[tag], base_zeta)
            idx.append([label, "zeta", tag]
                       + [ir[k] for k in ("n_total_idx", "n_senior_idx",
                                          "senior_share_pct", "cvar95_idx",
                                          "penalty_idx", "outsourced_idx",
                                          "release_idx")])
        book.save(RESULT_FILE)

    print("done: {0}".format(RESULT_FILE), flush=True)


if __name__ == "__main__":
    main()
