
import json
import math
import os
import time

import numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from Solvers import solver_benders as sb

_CACHE_INTEGER_STATUSES = {
    2: "ok",
    3: "infeasible",
    4: "inf_or_unbd",
    5: "unbounded",
    6: "cutoff",
    7: "iteration_limit",
    8: "node_limit",
    9: "time_limit",
    10: "solution_limit",
    11: "interrupted",
    12: "numeric",
    13: "suboptimal",
    15: "user_obj_limit",
    16: "work_limit",
    17: "oom",
}


def cache_identity(unit, instance):
    return str(unit), str(instance)


def load_cache(path):
    cache = {}
    if not os.path.exists(path):
        return cache
    with open(path, "r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                entry = json.loads(line)
            except (TypeError, ValueError):
                continue
            unit = entry.get("unit")
            instance = entry.get("instance")
            if isinstance(unit, str) and isinstance(instance, str) and "result" in entry:
                cache[cache_identity(unit, instance)] = entry["result"]
    return cache


def _cache_json_default(value):
    item = getattr(value, "item", None)
    if callable(item):
        return item()
    raise TypeError(repr(value))


def store_cache(path, unit, instance, result):
    cache_dir = os.path.dirname(path)
    if cache_dir and not os.path.exists(cache_dir):
        os.makedirs(cache_dir, exist_ok=True)

    payload = {"unit": str(unit), "instance": str(instance), "result": result}
    if os.path.exists(path) and os.path.getsize(path) > 0:
        with open(path, "rb+") as boundary:
            boundary.seek(-1, os.SEEK_END)
            if boundary.read(1) != b"\n":
                boundary.seek(0, os.SEEK_END)
                boundary.write(b"\n")
                boundary.flush()
                os.fsync(boundary.fileno())
    with open(path, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(
            payload,
            ensure_ascii=False,
            allow_nan=True,
            default=_cache_json_default,
        ) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def exception_status(error):
    text = "{0}: {1}".format(type(error).__name__, error).lower()
    oom_markers = (
        "out of memory", "out-of-memory", "memory limit", "mem_limit",
        "cannot allocate memory", "failed to allocate", "allocation failed",
        "gurobi error 10001",
    )
    if (isinstance(error, MemoryError)
            or getattr(error, "errno", None) == 10001
            or any(marker in text for marker in oom_markers)):
        return "oom"
    timeout_markers = ("time limit", "time-limit", "timed out", "timeout")
    if isinstance(error, TimeoutError) or any(
            marker in text for marker in timeout_markers):
        return "time_limit"
    return "error"


def error_text(error):
    return "{0}: {1}".format(type(error).__name__, error)


def normalized_solver_status(result, time_limit=None, time_field="time"):
    if not isinstance(result, dict):
        return "error"
    raw = result.get("status")
    status = None
    if isinstance(raw, bool):
        status = None
    elif isinstance(raw, (int, float)) and math.isfinite(float(raw)):
        status = _CACHE_INTEGER_STATUSES.get(
            int(raw), "status_{0}".format(int(raw)))
    elif raw is not None:
        status = str(raw).strip().lower().replace("-", "_").replace(" ", "_")
        aliases = {
            "optimal": "ok",
            "loaded": "ok",
            "timelimit": "time_limit",
            "memory_limit": "oom",
            "out_of_memory": "oom",
            "mem_limit": "oom",
        }
        status = aliases.get(status, status)
    elapsed = result.get(time_field)
    try:
        elapsed = float(elapsed)
    except (TypeError, ValueError):
        elapsed = None
    if status is None and time_limit is not None and elapsed is not None:
        threshold = max(
            float(time_limit) - max(1.0, 0.001 * float(time_limit)), 0.0)
        if elapsed >= threshold:
            status = "time_limit"
    if status is None:
        status = "ok" if result.get("objective") is not None else "no_solution"
    return status


def is_cacheable_result(result):
    return isinstance(result, dict) and result.get("_status", "ok") != "error"

INSTANCE_DIR = "instances"
SCALES = ("Small", "Medium")
INSTANCES_PER_SCALE = 5

RESULT_FILE = "results_joint_value.xlsx"
CACHE_FILE = "cache/cache_joint_value.jsonl"

MIP_GAP = 1e-4
TIME_LIMIT = 3600.0
THREADS = 0


def find_instances(directory, scale, count):
    found = []
    for name in sorted(os.listdir(directory)):
        if not name.endswith(".json"):
            continue
        stem = os.path.splitext(name)[0]
        parts = stem.rsplit("_", 1)
        if parts[0] != scale or not parts[1].isdigit():
            continue
        found.append((int(parts[1]), os.path.join(directory, name)))
    found.sort(key=lambda item: item[0])
    if len(found) < count:
        raise FileNotFoundError(
            "found {0} {1} instances but {2} required".format(
                len(found), scale, count))
    return found[:count]


def solve_variant(instance, lock_pref):
    started = time.perf_counter()
    try:
        res = sb.solve_instance(instance, {
            "mip_gap": MIP_GAP,
            "threads": THREADS,
            "time_limit": TIME_LIMIT,
            "extract_channels": True,
            "lock_pref": lock_pref,
        })
    except Exception as error:
        status = exception_status(error)
        if status not in ("oom", "time_limit"):
            raise
        return {
            "objective": None,
            "gap_pct": None,
            "wall_time_s": time.perf_counter() - started,
            "_status": status,
            "_error": error_text(error),
        }
    ch = res.get("channels") or {}
    out = {
        "objective": res.get("objective"),
        "gap_pct": res.get("gap"),
        "wall_time_s": res.get("time"),
        "F_commit": ch.get("F_commit"),
        "E_recourse": ch.get("E_recourse"),
        "cvar95": ch.get("cvar95"),
        "_status": normalized_solver_status(res, TIME_LIMIT, "time"),
    }
    if not lock_pref:
        out["mix_fraction"] = ch.get("mix_fraction")
    return out


def main():
    cache = load_cache(CACHE_FILE)
    book = Workbook()
    sheet = book.active
    sheet.title = "joint_value"
    headers = ["instance", "scale",
               "Z_joint", "Z_fixed", "value_pct",
               "mix_fraction_pct",
               "F_joint", "F_fixed",
               "EQ_joint", "EQ_fixed",
               "CVaR95_joint", "CVaR95_fixed",
               "time_joint_s", "time_fixed_s"]
    sheet.append(headers)
    for col, head in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = max(
            10, min(len(head) + 2, 20))
    book.save(RESULT_FILE)

    ordinal_prefix = {"Small": "S", "Medium": "M"}
    for scale in SCALES:
        selections = find_instances(INSTANCE_DIR, scale, INSTANCES_PER_SCALE)
        for ordinal, (_, path) in enumerate(selections, start=1):
            label = "{0}{1}".format(ordinal_prefix[scale], ordinal)
            with open(path, "r", encoding="utf-8") as handle:
                stored = json.load(handle)
            name = stored["meta"]["name"]

            variants = {}
            for lock in (False, True):
                unit = "lock={0}".format(int(lock))
                identity = cache_identity(unit, name)
                if identity in cache:
                    variants[lock] = cache[identity]
                    tag = "cache"
                else:
                    variants[lock] = solve_variant(stored, lock)
                    if is_cacheable_result(variants[lock]):
                        store_cache(CACHE_FILE, unit, name, variants[lock])
                        cache[identity] = variants[lock]
                    tag = "solved"
                print("[{0}] {1} lock={2} obj={3} ({4})".format(
                    label, name, lock,
                    variants[lock].get("objective"), tag), flush=True)

            joint = variants[False]
            fixed = variants[True]
            z_joint = joint.get("objective")
            z_fixed = fixed.get("objective")
            value_pct = (100.0 * (z_fixed - z_joint) / z_fixed
                         if z_joint is not None and z_fixed not in (None, 0.0)
                         else None)
            mix_pct = (100.0 * joint.get("mix_fraction", float("nan"))
                       if joint.get("mix_fraction") is not None else None)
            sheet.append([
                label, scale,
                z_joint, z_fixed, value_pct, mix_pct,
                joint.get("F_commit"), fixed.get("F_commit"),
                joint.get("E_recourse"), fixed.get("E_recourse"),
                joint.get("cvar95"), fixed.get("cvar95"),
                joint.get("wall_time_s"), fixed.get("wall_time_s"),
            ])
            book.save(RESULT_FILE)
            print("[{0}] value={1}%  mix={2}%".format(
                label,
                "nan" if value_pct is None else "{0:.3f}".format(value_pct),
                "nan" if mix_pct is None else "{0:.3f}".format(mix_pct)),
                flush=True)
    print("done: {0}".format(RESULT_FILE), flush=True)


if __name__ == "__main__":
    main()
